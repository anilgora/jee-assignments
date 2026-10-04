#!/usr/bin/env python3
"""Build the calibrated difficulty report (Excel) from a ratings JSON file.

Usage:
    python build_report.py ratings.json output.xlsx

Input schema
------------
{
  "source": "Mock_Test_7.pdf",            # file name(s) rated
  "evaluated_on": "2026-09-26",           # optional
  "topic_adjustment": 0.5,                # optional, default 0.5 (equivalence relations/classes)
  "mcq_multi_concept_floor": 2.5,         # optional, default 2.5
  "questions": [
    {
      "sr_no": 1,                          # int, sequential from 1
      "paper": "P1",                       # paper/part label; use "" for a single paper
      "q_ref": "Q12",                      # as printed, e.g. "Q12(ii)"
      "type": "Numerical",                 # SCQ | MCQ | Numerical | INT | MTC SCQ | Q Stem | Para SCQ
      "chapter": "Definite Integration",
      "marks": 4,                          # full marks for the question
      "level": 4,                          # int 0-5: the user's-scale level (format already included)
      "eq_relation": false,                # optional: key step tests equivalence relations/classes
      "multi_concept": false,              # optional: two or more distinct concepts involved
      "key_step": "Spot the king-type substitution.",
      "justification": "Two to three sentences naming what sets the level.",
      "nearest_anchor": "2025-P2-Q16 (L4)",
      "confidence": "High"                 # High | Medium | Low
    }
  ]
}

All fields except eq_relation and multi_concept are required. The script stops with a clear message on a missing field instead of
writing a half-built workbook.

Sheets produced
---------------
Evaluation  one row per question: level, equivalence-relation adjustment, MCQ floor, adjusted
            level (formulas), marks x adjusted level
Summary     level distribution; weighted index (adjusted and unadjusted) per paper and overall
Reference   the user's own weighted index for official JEE Advanced 2021-2026, from anchors.csv
            (the user's ratings already include everything, so no adjustment is applied to them)

Adjustments (user's instructions, 27-Sep-2026)
----------------------------------------------
Adjusted level = min(5, max(level + topic adj, floor))
  topic adj: +0.5 when eq_relation is true (equivalence relations / classes are new to the
             JEE Advanced syllabus and students have little practice material)
  floor    : 2.5 for an MCQ (one or more correct) with multi_concept true; otherwise none
No format adjustment: the level itself already reflects the question format.

The workbook uses live formulas. Run the xlsx skill's recalc.py on the output (if available) so
cached values exist for previewers.
"""

import csv
import json
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LEVEL_LABELS = {
    0: "Easier than JEE Main",
    1: "JEE Main level",
    2: "Easier than Advanced",
    3: "Average Advanced",
    4: "Slightly difficult",
    5: "Very difficult",
}
LEVEL_FILLS = {0: "D9EAD3", 1: "E8F0D8", 2: "FFF2CC", 3: "FCE5CD", 4: "F4CCCC", 5: "EA9999"}

DEFAULT_TOPIC_ADJ = 0.5
DEFAULT_MCQ_FLOOR = 2.5
ANCHORS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "references", "anchors.csv")


def reference_rows():
    """User's own ratings of official papers, as they are (they already include every factor)."""
    by_year = {}
    with open(ANCHORS, encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            lv, mk = int(r["Level"]), int(r["Marks"])
            y = by_year.setdefault(int(r["Year"]), [0, 0, 0, 0, 0, 0])
            y[0] += mk; y[1] += mk * lv; y[2] += lv; y[3] += 1
            y[4] += mk if lv <= 1 else 0; y[5] += mk if lv >= 4 else 0
            y.append(lv == 5)
    return [(yr, v[1] / v[0], v[2] / v[3], v[4], v[5], sum(v[6:])) for yr, v in sorted(by_year.items())]


REQUIRED = ["sr_no", "paper", "q_ref", "type", "chapter", "marks", "level",
            "key_step", "justification", "nearest_anchor", "confidence"]

HEADER_FILL = PatternFill("solid", fgColor="1A365D")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=13, bold=True, color="1A365D")
NOTE = Font(name="Arial", size=9, italic=True, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top")


def fail(msg):
    sys.exit(f"build_report.py: {msg}")


def header(ws, row, cols):
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width


def main(src, out):
    data = json.load(open(src, encoding="utf-8"))
    qs = data.get("questions") or fail("no questions in JSON")
    for q in qs:
        missing = [k for k in REQUIRED if k not in q]
        if missing:
            fail(f"question {q.get('sr_no', '?')} is missing {missing}")
        if q["level"] not in LEVEL_LABELS:
            fail(f"question {q['sr_no']} has level {q['level']}; must be 0-5")

    wb = Workbook()

    # ---------------- Evaluation ----------------
    ws = wb.active
    ws.title = "Evaluation"
    topic_adj = float(data.get("topic_adjustment", DEFAULT_TOPIC_ADJ))
    mcq_floor = float(data.get("mcq_multi_concept_floor", DEFAULT_MCQ_FLOOR))
    cols = [("Sr No", 7), ("Paper", 8), ("Q Ref", 9), ("Type", 12), ("Chapter", 24),
            ("Marks", 7), ("Level", 7), ("Level Label", 20), ("Equiv. Rel. Adj", 9),
            ("MCQ Floor", 8), ("Adjusted Level", 9), ("Marks x Adj Level", 10),
            ("Key Step", 38), ("Justification", 60), ("Nearest Anchor", 20), ("Confidence", 11)]
    header(ws, 1, cols)
    for r, q in enumerate(qs, 2):
        ta = topic_adj if q.get("eq_relation") else 0
        fl = mcq_floor if (q["type"] == "MCQ" and q.get("multi_concept")) else 0
        vals = [q["sr_no"], q["paper"], q["q_ref"], q["type"], q["chapter"], q["marks"],
                q["level"], LEVEL_LABELS[q["level"]], ta, fl, f"=MIN(5,MAX(G{r}+I{r},J{r}))",
                f"=F{r}*K{r}", q["key_step"], q["justification"], q["nearest_anchor"],
                q["confidence"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = BODY, BORDER
            cell.alignment = WRAP if c in (5, 13, 14, 15) else CENTER
        for c in (9, 10, 11, 12):
            ws.cell(row=r, column=c).number_format = "0.0"
        for c in (7, 8):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=LEVEL_FILLS[q["level"]])
    last = len(qs) + 1
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:P{last}"
    ws.cell(row=last + 2, column=1, value=(
        f"Adjusted level = min(5, max(level + equivalence-relation adj, MCQ floor)). Equivalence adj "
        f"+{topic_adj} when the key step tests equivalence relations/classes; MCQ floor {mcq_floor} for "
        f"one-or-more-correct questions involving several concepts. Edit columns I-J to change them.")).font = NOTE

    rng = lambda col: f"Evaluation!${col}$2:${col}${last}"

    # ---------------- Summary ----------------
    sm = wb.create_sheet("Summary")
    sm["A1"] = "Difficulty summary"
    sm["A1"].font = TITLE
    sm["A2"] = f"Source: {data.get('source', '')}   |   Evaluated on: {data.get('evaluated_on', str(date.today()))}"
    sm["A2"].font = NOTE

    header(sm, 4, [("Level", 8), ("Label", 22), ("Questions", 11), ("Marks", 9), ("% of marks", 11)])
    for i, lv in enumerate(range(6)):
        r = 5 + i
        row = [lv, LEVEL_LABELS[lv], f"=COUNTIF({rng('G')},A{r})",
               f"=SUMIF({rng('G')},A{r},{rng('F')})", f"=IF(SUM($D$5:$D$10)=0,0,D{r}/SUM($D$5:$D$10))"]
        for c, v in enumerate(row, 1):
            cell = sm.cell(row=r, column=c, value=v)
            cell.font, cell.border = BODY, BORDER
            cell.fill = PatternFill("solid", fgColor=LEVEL_FILLS[lv])
        sm.cell(row=r, column=5).number_format = "0%"
    sm["A11"], sm["C11"], sm["D11"] = "Total", "=SUM(C5:C10)", "=SUM(D5:D10)"
    for c in ("A11", "C11", "D11"):
        sm[c].font, sm[c].border = BOLD, BORDER

    papers = sorted({q["paper"] for q in qs if q["paper"]})
    header(sm, 13, [("Scope", 8), ("Weighted index (adjusted)", 22), ("Weighted index (unadjusted)", 11),
                    ("Marks", 9), ("Mean level", 11)])
    r = 14
    for p in papers:
        sm.cell(row=r, column=1, value=p)
        sm.cell(row=r, column=2, value=f"=SUMIF({rng('B')},A{r},{rng('L')})/SUMIF({rng('B')},A{r},{rng('F')})")
        sm.cell(row=r, column=3, value=f"=SUMPRODUCT(({rng('B')}=A{r})*{rng('F')}*{rng('G')})/SUMIF({rng('B')},A{r},{rng('F')})")
        sm.cell(row=r, column=4, value=f"=SUMIF({rng('B')},A{r},{rng('F')})")
        sm.cell(row=r, column=5, value=f"=AVERAGEIF({rng('B')},A{r},{rng('G')})")
        r += 1
    sm.cell(row=r, column=1, value="Overall")
    sm.cell(row=r, column=2, value=f"=SUM({rng('L')})/SUM({rng('F')})")
    sm.cell(row=r, column=3, value=f"=SUMPRODUCT({rng('F')},{rng('G')})/SUM({rng('F')})")
    sm.cell(row=r, column=4, value=f"=SUM({rng('F')})")
    sm.cell(row=r, column=5, value=f"=AVERAGE({rng('G')})")
    overall_row = r
    for rr in range(14, r + 1):
        for c in range(1, 6):
            cell = sm.cell(row=rr, column=c)
            cell.font, cell.border = (BOLD if rr == r else BODY), BORDER
        for c in (2, 3, 5):
            sm.cell(row=rr, column=c).number_format = "0.00"
    sm.cell(row=r + 2, column=1, value=(
        "Weighted index = sum(marks x level) / sum(marks), scale 0-5. 'Adjusted' uses the adjusted "
        "level (equivalence-relation adjustment and MCQ floor); 'unadjusted' uses the integer level.")).font = NOTE

    # ---------------- Reference ----------------
    rf = wb.create_sheet("Reference")
    rf["A1"] = "User's own ratings of official JEE Advanced maths papers"
    rf["A1"].font = TITLE
    header(rf, 3, [("Year", 12), ("Weighted index", 15), ("Mean level", 11),
                   ("Marks at L0-1", 13), ("Marks at L4-5", 13), ("Level-5 questions", 15)])
    ref = reference_rows()
    for i, row in enumerate(ref):
        for c, v in enumerate(row, 1):
            cell = rf.cell(row=4 + i, column=c, value=v)
            cell.font, cell.border = BODY, BORDER
    r = 4 + len(ref)
    rf.cell(row=r, column=1, value="This paper")
    rf.cell(row=r, column=2, value=f"=Summary!B{overall_row}")
    rf.cell(row=r, column=3, value=f"=Summary!E{overall_row}")
    rf.cell(row=r, column=4, value="=Summary!D5+Summary!D6")
    rf.cell(row=r, column=5, value="=Summary!D9+Summary!D10")
    rf.cell(row=r, column=6, value="=Summary!C10")
    for c in range(1, 7):
        cell = rf.cell(row=r, column=c)
        cell.font, cell.border = BOLD, BORDER
    for rr in range(4, r + 1):
        rf.cell(row=rr, column=2).number_format = "0.00"
        rf.cell(row=rr, column=3).number_format = "0.00"
    rf.cell(row=r + 2, column=1, value=(
        "Source: the user's ratings in references/anchors.csv, used as given (they already reflect "
        "format and topic). 'This paper' uses the adjusted index. Official papers are 120 marks; "
        "compare the marks-at-level columns only for full-length papers.")).font = NOTE

    wb.save(out)
    print(f"Wrote {out}: {len(qs)} questions, papers={papers or ['(single)']}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        fail("usage: python build_report.py ratings.json output.xlsx")
    main(sys.argv[1], sys.argv[2])
