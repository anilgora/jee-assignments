#!/usr/bin/env python3
"""Step 2 and Step 3: write the blueprint workbook, and re-check it after the user edits it.

Build:  python make_blueprint.py build --config input_config.xlsx --workdir WORKDIR --picks picks.json \
            --paper-name "CD PRT 04 Paper 01" --out Question_Paper_Blueprint.xlsx
Check:  python make_blueprint.py check --config input_config.xlsx --blueprint Question_Paper_Blueprint.xlsx \
            --pattern main|advanced [--json WORKDIR/blueprint.json]

picks.json: [{"sr": 1, "concept1": "...", "concept2": "..." | null, "concept3": "..." | null,
              "chapter2": optional, "chapter3": optional, "link": "one line: why these concepts link"}]
chapter2 / chapter3 are looked up from the Concepts sheet when omitted. Concept 1 belongs to the slot's
chapter. Concepts 2 and 3 must belong to chapters listed in the Chapters sheet. No concept is used twice.
Main pattern: at most two concepts per question and a 20/40/40 mix (single / same chapter / different
chapters), each within +-1 of target. Advanced pattern: two or three concepts per question, at least 2 questions
with three concepts, and at least 40% of the questions mixing different chapters.
--exclude OLDEST.xlsx ... NEWEST.xlsx: blueprints of earlier papers of the same phase. A concept from them (or
repeated inside this paper) is allowed only when its chapter has no fresh concept left, and then the earliest-used
ones are reused first; repeated concept combinations are warned about. Errors on build, warnings on check.
"""
import argparse, json, math, os, re, sys
import pandas as pd
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import *

COLS = ['Sr No', 'Chapter', 'Concept1', 'Concept 2', 'Concept 3', 'Chapter of Concept 2', 'Chapter of Concept 3',
        'Pairing Type', 'Link Note', 'Q Type', 'Difficulty Level', 'Answer Plan', 'Word File and Folder name']


def concept_index(cfg):
    idx = {}
    for ch, lst in cfg['concepts'].items():
        for c in lst:
            idx.setdefault(c, []).append(ch)
    return idx


def pairing_type(ch1, c2, ch2, c3, ch3):
    if not c2 and not c3:
        return 'Single concept'
    chs = [x for x in (ch2, ch3) if x]
    if all(same_chapter(ch1, x) for x in chs):
        return 'Same chapter'
    return 'Different chapters'


def validate_rows(rows, cfg, pattern, excluded=None, strict=False, excluded_pairs=None):
    """rows: dicts with sr, chapter, c1, c2, c3, ch2, ch3, q_type, difficulty. Returns (errors, warnings, kinds).
    strict=True (blueprint build) turns the paper-rule shortfalls into errors; check mode only warns."""
    from collections import Counter
    errs, warns = [], []
    rule = errs if strict else warns
    excluded = excluded or {}
    idx = concept_index(cfg)
    chapters = list(cfg['chapters'])
    uses = Counter()
    for r in rows:
        t = f"Q{r['sr']}"
        for k, c, ch in (('Concept1', r['c1'], r['chapter']), ('Concept 2', r['c2'], r['ch2']), ('Concept 3', r['c3'], r['ch3'])):
            if not c:
                continue
            uses[c] += 1
            if c not in idx:
                errs.append(f'{t}: {k} "{c}" is not in the Concepts sheet.' + suggest(c, list(idx)))
                continue
            if ch not in idx[c]:
                errs.append(f'{t}: {k} "{c}" belongs to {idx[c]}, not "{ch}".')
            if k != 'Concept1' and ch not in chapters:
                errs.append(f'{t}: {k} chapter "{ch}" is not in the paper\'s Chapters sheet.')
        if r['c3'] and not r['c2']:
            errs.append(f'{t}: Concept 3 given without Concept 2.')
        if pattern == 'main' and r['c3']:
            errs.append(f'{t}: main pattern allows at most two concepts.')
        if r['c1'] and r['c2'] and r['c1'] == r['c2']:
            errs.append(f'{t}: Concept 2 equals Concept 1.')
        if type_code(r['q_type']) is None:
            errs.append(f'{t}: unsupported question type "{r["q_type"]}".')
        if canon_level(r['difficulty']) is None:
            errs.append(f'{t}: unknown difficulty "{r["difficulty"]}".')
    # ---- concept reuse. A chapter's concept may be reused only when it has no fresh concept left
    # (fresh = not used in an earlier paper of the phase and not already used in this paper).
    fresh = {ch: [c for c in lst if c not in excluded and uses.get(c, 0) == 0] for ch, lst in cfg['concepts'].items()}
    seen = set()
    for r in rows:
        for c in (r['c1'], r['c2'], r['c3']):
            if not c or c not in idx or (c, 'x') in seen:
                continue
            seen.add((c, 'x'))
            ch = idx[c][0]
            if uses[c] > 1:
                msg = f'concept "{c}" is used {uses[c]} times in this paper'
                if fresh.get(ch):
                    rule.append(f'{msg}, but chapter "{ch}" still has unused concepts, e.g. "{fresh[ch][0]}".')
                else:
                    warns.append(f'{msg} (chapter "{ch}" has too few concepts; allowed, but do not repeat the same pair).')
            if c in excluded:
                if fresh.get(ch):
                    rule.append(f'concept "{c}" was used in {excluded[c][0]}, but chapter "{ch}" still has fresh concepts, e.g. "{fresh[ch][0]}".')
    # reused concepts of a chapter should be the earliest-used ones
    by_ch = {}
    for c in {c for c in uses if c in excluded and c in idx}:
        by_ch.setdefault(idx[c][0], []).append(c)
    for ch, reused in by_ch.items():
        old = sorted([c for c in cfg['concepts'][ch] if c in excluded], key=lambda c: excluded[c][1])
        if set(reused) != set(old[:len(reused)]):
            warns.append(f'chapter "{ch}": reuse the earliest-used concepts first (earliest: {old[:len(reused)]}).')
    # ---- repeated pairs (avoid when possible)
    pairs = Counter()
    for r in rows:
        cs = frozenset(x for x in (r['c1'], r['c2'], r['c3']) if x)
        if len(cs) >= 2:
            pairs[cs] += 1
            if excluded_pairs and cs in excluded_pairs:
                warns.append(f"Q{r['sr']}: this concept combination was already used in {excluded_pairs[cs]}; change it if possible.")
    for cs, k in pairs.items():
        if k > 1:
            warns.append(f'the combination {sorted(cs)} appears {k} times in this paper; change it if possible.')
    # ---- distribution against config (case-insensitive label match)
    for key, col in (('chapters', 'chapter'), ('difficulties', 'difficulty'), ('qtypes', 'q_type')):
        have = Counter(str(r[col]).strip().lower() for r in rows)
        want = {k.lower(): v for k, v in cfg[key].items()}
        for k in sorted(set(have) | set(want)):
            if have.get(k, 0) != want.get(k, 0):
                warns.append(f'{key}: "{k}" has {have.get(k, 0)} rows, config says {want.get(k, 0)}.')
    ranks = [type_rank(type_code(r['q_type']) or '', pattern) for r in rows]
    if any(x > y for x, y in zip(ranks, ranks[1:])):
        want = ' > '.join(ORDER[pattern])
        rule.append(f'Question order is not by type. Order for the {pattern} pattern: {want}.')
    n = len(rows)
    kinds = Counter(pairing_type(r['chapter'], r['c2'], r['ch2'], r['c3'], r['ch3']) for r in rows)
    if pattern == 'advanced':
        need = math.ceil(0.4 * n)
        if kinds.get('Different chapters', 0) < need:
            rule.append(f'Advanced needs at least {need} of {n} questions mixing different chapters; found {kinds.get("Different chapters", 0)}.')
        for r in rows:
            if not r['c2']:
                rule.append(f"Q{r['sr']}: Advanced questions must have two or three concepts.")
        three = sum(1 for r in rows if r['c3'])
        if three < 2:
            rule.append(f'Advanced needs at least 2 questions with three concepts; found {three}.')
    if pattern == 'main':
        s = round(0.2 * n); sm = round(0.4 * n); cr = n - s - sm
        for name, want in (('Single concept', s), ('Same chapter', sm), ('Different chapters', cr)):
            if abs(kinds.get(name, 0) - want) > 1:
                warns.append(f'Mix: {name} = {kinds.get(name, 0)}, target {want} (+-1 allowed).')
    return errs, warns, kinds


def build(a):
    cfg = load_config(a.config)
    errs, _ = check_config(cfg)
    if errs:
        sys.exit('ERROR in input_config:\n  ' + '\n  '.join(errs))
    plan = json.load(open(os.path.join(a.workdir, 'slots.json')))
    picks = {p['sr']: p for p in json.load(open(a.picks))}
    idx = concept_index(cfg)
    rows = []
    for s in plan['slots']:
        p = picks.get(s['sr'])
        if not p:
            sys.exit(f"No pick for Q{s['sr']}.")
        def chap(c, given):
            if not c: return ''
            if given: return given
            ch = idx.get(c, [])
            return ch[0] if ch else ''
        c2 = (p.get('concept2') or '').strip(); c3 = (p.get('concept3') or '').strip()
        c1 = p['concept1'].strip()
        ch2 = chap(c2, p.get('chapter2')); ch3 = chap(c3, p.get('chapter3'))
        rows.append({'sr': s['sr'], 'chapter': s['chapter'], 'c1': c1, 'c2': c2, 'c3': c3, 'ch2': ch2, 'ch3': ch3,
                     'q_type': s['q_type'], 'difficulty': s['difficulty'], 'link': p.get('link', ''),
                     'answer_plan': s.get('answer_plan', '')})
    excluded = used_concepts(a.exclude)
    errs, warns, kinds = validate_rows(rows, cfg, plan['pattern'], excluded, strict=True, excluded_pairs=used_pairs(a.exclude))
    if errs:
        sys.exit('Blueprint NOT written. Fix these:\n  ' + '\n  '.join(errs))
    out = []
    for r in rows:
        out.append([r['sr'], r['chapter'], r['c1'], r['c2'], r['c3'], r['ch2'], r['ch3'],
                    pairing_type(r['chapter'], r['c2'], r['ch2'], r['c3'], r['ch3']), r['link'], r['q_type'], r['difficulty'], r['answer_plan'], a.paper_name])
    pd.DataFrame(out, columns=COLS).to_excel(a.out, index=False)
    fmt(a.out)
    print(f'Wrote {a.out}: {len(rows)} questions. Mix: {dict(kinds)}')
    for w in warns: print('WARN', w)


def fmt(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = load_workbook(path); ws = wb.active
    widths = [6, 24, 44, 44, 38, 24, 24, 18, 60, 30, 34, 34, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'C2'
    wb.save(path)


ALIASES = {
    'sr': ['sr no', 'sr', 'sr.', 'sr. no', 'q no', 'question no', 'no', 'no.', 's no', 's.no'],
    'chapter': ['chapter', 'chapter 1', 'chapter1'],
    'c1': ['concept1', 'concept 1', 'main concept'],
    'c2': ['concept 2', 'concept2'],
    'c3': ['concept 3', 'concept3'],
    'ch2': ['chapter of concept 2', 'chapter 2', 'chapter2'],
    'ch3': ['chapter of concept 3', 'chapter 3', 'chapter3'],
    'q_type': ['q type', 'question type', 'type', 'qtype'],
    'difficulty': ['difficulty level', 'difficulty', 'level'],
    'link': ['link note', 'link'],
    'answer_plan': ['answer plan', 'answer'],
    'paper': ['word file and folder name', 'paper name', 'paper', 'folder name'],
}


def _split_cell(v, known=None, notes=None, where=''):
    """A cell may hold several concepts, one per LINE. Concept names themselves contain ';', ',' and '|', so those are never
    split blindly. With the Concepts sheet (`known`), a '; '-separated list is split only if every piece is a known concept."""
    parts = [p.strip().rstrip(';').strip() for p in re.split(r'\r?\n', str(v)) if p.strip().rstrip(';').strip()]
    out = []
    for p in parts:
        if known and p not in known and ';' in p:
            pieces = [x.strip() for x in p.split(';') if x.strip()]
            if len(pieces) > 1 and all(x in known for x in pieces):
                out.extend(pieces)
                continue
        if not known and ';' in p and notes is not None:
            notes.append(f'{where}: a concept cell contains ";" and was kept as ONE concept (concept names can contain ";"). To list several concepts, use separate columns (Concept 1, Concept 2, ...) or one concept per line.')
        out.append(p)
    return out


def read_blueprint(path, known=None):
    """Read a blueprint workbook with tolerant column names and ANY number of concepts per question.
    Concepts come from columns 'Concept 1', 'Concept 2', ... (any count), or from a 'Concepts' column, or from several
    concepts in one cell, one per line. Returns (rows, missing_required, notes)."""
    df = pd.read_excel(path).fillna('')
    low = {str(c).strip().lower(): c for c in df.columns}
    col = {k: next((low[n] for n in names if n in low), None) for k, names in ALIASES.items() if k not in ('c1', 'c2', 'c3', 'ch2', 'ch3')}
    ccols = []        # (n, column name)
    for name, orig in low.items():
        m = re.fullmatch(r'(?:concept|main concept)\s*(\d*)', name)
        if m:
            ccols.append((int(m.group(1) or 1), orig))
    ccols.sort()
    list_col = next((low[n] for n in ('concepts', 'concept list', 'concept(s)') if n in low), None)
    chcols = {}
    for name, orig in low.items():
        m = re.fullmatch(r'chapter\s*(?:of\s*concept\s*)?(\d+)', name)
        if m:
            chcols[int(m.group(1))] = orig
    miss = [k for k in ('chapter', 'q_type', 'difficulty') if col[k] is None]
    if not ccols and list_col is None:
        miss.append('c1')
    rows = []
    notes = []
    for i, (_, r) in enumerate(df.iterrows(), 1):
        g = lambda k: str(r[col[k]]).strip() if col.get(k) is not None else ''
        cs, chs = [], []
        for n, cname in ccols:
            parts = _split_cell(r[cname], known, notes, f'row {i}')
            ch_n = str(r[chcols[n]]).strip() if n in chcols and n > 1 else ''
            for p in parts:
                cs.append(p); chs.append(ch_n if len(parts) == 1 else '')
        if list_col is not None:
            for p in _split_cell(r[list_col], known, notes, f'row {i}'):
                cs.append(p); chs.append('')
        try:
            sr = int(float(g('sr')))
        except ValueError:
            sr = i
        if not (cs or g('q_type') or g('chapter')):
            continue
        rows.append({'sr': sr, 'chapter': g('chapter'), 'cs': cs, 'chs': chs,
                     'c1': cs[0] if cs else '', 'c2': cs[1] if len(cs) > 1 else '', 'c3': cs[2] if len(cs) > 2 else '',
                     'ch2': chs[1] if len(chs) > 1 else '', 'ch3': chs[2] if len(chs) > 2 else '',
                     'q_type': g('q_type'), 'difficulty': g('difficulty'), 'link': g('link'), 'answer_plan': g('answer_plan'),
                     'paper': g('paper')})
    return rows, miss, notes


def detect_pattern_rows(rows):
    codes = {type_code(r['q_type']) for r in rows}
    levels = [LEVEL_OF[canon_level(r['difficulty'])] for r in rows if canon_level(r['difficulty'])]
    adv = sorted(c for c in codes & ADV_ONLY_TYPES if c)
    hard = any(l >= 4 for l in levels)
    if adv or hard:
        why = []
        if adv: why.append('Advanced-only question types present: ' + ', '.join(adv))
        if hard: why.append('questions at level 4-5 present')
        return 'advanced', why
    if codes <= {'SCQ', 'Numerical'}:
        return 'main', ['only Single Correct and Numerical types, no question above level 3']
    return None, ['could not tell']


def validate_direct(rows, pattern, excluded=None, excluded_pairs=None):
    """Checks that need no input_config: used when the user supplies only a blueprint. Any number of concepts per row."""
    from collections import Counter
    errs, warns = [], []
    seen = {}
    for r in rows:
        t = f"Q{r['sr']}"
        cs = r['cs']
        if not cs:
            errs.append(f'{t}: no concept given.')
        if type_code(r['q_type']) is None:
            errs.append(f'{t}: unsupported question type "{r["q_type"]}". Known: {list(TYPE_OF)}')
        if canon_level(r['difficulty']) is None:
            errs.append(f'{t}: unknown difficulty "{r["difficulty"]}". Known: {list(LEVEL_OF)}')
        if pattern == 'main' and len(cs) > 2:
            warns.append(f'{t}: the Main pattern normally has at most two concepts; this blueprint lists {len(cs)}. Kept as given: the question will use all {len(cs)}.')
        if pattern == 'advanced' and len(cs) < 2:
            warns.append(f'{t}: Advanced questions should have two or more concepts; this one has {len(cs)}.')
        for c in cs:
            if c in seen:
                warns.append(f'concept "{c}" is used in {seen[c]} and {t}.')
            seen[c] = t
            if excluded and c in excluded:
                warns.append(f'{t}: concept "{c}" was used in an earlier paper of this phase ({excluded[c][0]}).')
    if pattern == 'advanced':
        three = sum(1 for r in rows if len(r['cs']) >= 3)
        if three < 2:
            warns.append(f'Advanced papers should have at least 2 questions with three or more concepts; found {three}.')
        if any(any(r['chs'][1:]) for r in rows):
            cross = sum(1 for r in rows if any(ch and not same_chapter(r['chapter'], ch) for ch in r['chs'][1:]))
            need = math.ceil(0.4 * len(rows))
            if cross < need:
                warns.append(f'Advanced papers should have at least {need} of {len(rows)} questions mixing chapters; found {cross}.')
        else:
            warns.append('No chapter column for the extra concepts: the 40% chapter-mixing rule could not be checked.')
    pairs = Counter()
    for r in rows:
        cs = frozenset(r['cs'])
        if len(cs) >= 2:
            pairs[cs] += 1
            if excluded_pairs and cs in excluded_pairs:
                warns.append(f"Q{r['sr']}: this concept combination was already used in {excluded_pairs[cs]}.")
    for cs, k in pairs.items():
        if k > 1:
            warns.append(f'the combination {sorted(cs)} appears {k} times.')
    return errs, warns


def check(a):
    direct = not a.config
    known = None
    if a.config:
        _cfg = load_config(a.config)
        known = {c for lst in _cfg['concepts'].values() for c in lst}
    rows, miss, read_notes = read_blueprint(a.blueprint, known)
    if miss:
        sys.exit(f'Blueprint is missing required columns (looked for aliases of): {miss}')
    pattern = a.pattern
    if pattern == 'auto':
        pattern, why = detect_pattern_rows(rows)
        if pattern is None:
            sys.exit('Cannot tell Main from Advanced from this blueprint (' + '; '.join(why) + '). Ask the user and rerun with --pattern.')
        print(f'Pattern inferred: {pattern.upper()} ({"; ".join(why)}). Confirm with the user if the paper name or phase says otherwise.')
    notes = []
    # question order by type (user's rule); --reorder applies it and prints the mapping
    ranks = [type_rank(type_code(r['q_type']) or '', pattern) for r in rows]
    if any(x > y for x, y in zip(ranks, ranks[1:])):
        if a.reorder:
            order = sorted(range(len(rows)), key=lambda i: (ranks[i], i))
            old = [rows[i]['sr'] for i in order]
            rows = [rows[i] for i in order]
            for new_sr, r in enumerate(rows, 1):
                r['old_sr'] = r['sr']; r['sr'] = new_sr
            mapping = ', '.join(f"old Q{o}->new Q{n}" for n, o in enumerate(old, 1) if o != n)
            notes.append(f'Questions reordered by type for the {pattern} pattern: {mapping}')
        else:
            notes.append(f'Question order is not by type (order for {pattern}: ' + ' > '.join(ORDER[pattern]) + '). Use --reorder to apply it.')
    excluded = used_concepts(a.exclude)
    if direct:
        errs, warns = validate_direct(rows, pattern, excluded, used_pairs(a.exclude))
        kinds = {}
    else:
        cfg = load_config(a.config)
        idx = concept_index(cfg)
        for r in rows:
            if r['c2'] and not r['ch2']: r['ch2'] = idx.get(r['c2'], [''])[0]
            if r['c3'] and not r['ch3']: r['ch3'] = idx.get(r['c3'], [''])[0]
        errs, warns, kinds = validate_rows(rows, cfg, pattern, excluded, excluded_pairs=used_pairs(a.exclude))
        warns = [w for w in warns if not w.startswith('Question order is not by type')]
        for r in rows:
            if len(r['cs']) > 3:
                warns.append(f"Q{r['sr']}: {len(r['cs'])} concepts listed; concepts after the third were not checked against the Concepts sheet, but the question must still use all of them.")
    # answer plans: validate what the user left, draw plans for rows that have none
    import random as _r
    codes = [type_code(r['q_type']) for r in rows]
    fresh = make_answer_plans(codes, _r.Random(a.seed))
    for r, c, f in zip(rows, codes, fresh):
        if not r['answer_plan']:
            r['answer_plan'] = f
            warns.append(f"Q{r['sr']}: no Answer Plan in the blueprint; planned '{f}'.")
        elif any(p['kind'] == 'bad' for p in parse_plan(r['answer_plan'])):
            errs.append(f"Q{r['sr']}: cannot read Answer Plan '{r['answer_plan']}'. Use 'Option C', 'Options B, D', 'Q1: Option A; Q2: Option C', 'Integer 7' or a numerical style.")
    print(f'{len(rows)} questions ({"direct blueprint, no input_config" if direct else "checked against input_config"}).' + (f' Mix: {dict(kinds)}' if kinds else ''))
    for n in notes: print('NOTE', n)
    for n in read_notes: print('WARN', n)
    for w in warns: print('WARN', w)
    for e in errs: print('ERROR', e)
    if a.json:
        out = []
        for r in rows:
            out.append({'sr': r['sr'], 'old_sr': r.get('old_sr', r['sr']), 'chapter': r['chapter'],
                        'concepts': r['cs'], 'n_concepts': len(r['cs']),
                        'concept_chapters': [r['chapter']] + [x for x in r['chs'][1:] if x],
                        'pairing': ('Single concept' if len(r['cs']) <= 1 else ('unknown' if not any(r['chs'][1:]) else ('Different chapters' if any(ch and not same_chapter(r['chapter'], ch) for ch in r['chs'][1:]) else 'Same chapter'))),
                        'link': r['link'], 'q_type': r['q_type'], 'type_code': type_code(r['q_type']), 'difficulty': r['difficulty'],
                        'target_level': LEVEL_OF[canon_level(r['difficulty'])] if canon_level(r['difficulty']) else None,
                        'answer_plan': r['answer_plan'],
                        'paper': r['paper'] or os.path.splitext(os.path.basename(a.blueprint))[0], 'pattern': pattern})
        json.dump(out, open(a.json, 'w'), indent=1, ensure_ascii=False)
        print('Wrote', a.json)
    sys.exit(1 if errs else 0)


if __name__ == '__main__':
    ap = argparse.ArgumentParser(); sub = ap.add_subparsers(dest='cmd', required=True)
    b = sub.add_parser('build'); b.add_argument('--config', required=True); b.add_argument('--workdir', required=True)
    b.add_argument('--picks', required=True); b.add_argument('--exclude', nargs='*', default=[]); b.add_argument('--paper-name', default='Paper 01'); b.add_argument('--out', default='Question_Paper_Blueprint.xlsx')
    c = sub.add_parser('check'); c.add_argument('--config', default=None); c.add_argument('--blueprint', required=True)
    c.add_argument('--pattern', choices=['main', 'advanced', 'auto'], default='auto'); c.add_argument('--reorder', action='store_true'); c.add_argument('--seed', type=int, default=0); c.add_argument('--json'); c.add_argument('--exclude', nargs='*', default=[])
    a = ap.parse_args()
    build(a) if a.cmd == 'build' else check(a)
