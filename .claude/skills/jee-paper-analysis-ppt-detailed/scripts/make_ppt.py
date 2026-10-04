#!/usr/bin/env python3
"""
make_ppt.py (DETAILED-SOLUTIONS edition) - build a "Paper Analysis" deck from a content JSON file.

Usage (from a terminal, in the folder that holds the JSON):
    pip install python-pptx pypandoc_binary        # once
    python make_ppt.py PRT_06_Paper_01_content.json
    python make_ppt.py PRT_06_Paper_01_content.json -o "PRT 06 Analysis.pptx" --max-bullets 5

Same deck as the standard edition (title, section dividers, question slides with level badge,
difficulty table, summary, difficulty profile, comparison with JEE Advanced papers) EXCEPT for
the solution part, which is:
    1. PRIMARY SOLUTION - the COMPLETE solution from the PDF (every step), verified and
       corrected where the PDF is wrong. Opens with the amber "Key idea" strip, ends with the
       green "Correct Answer" box. Speaker notes say whether the PDF solution was kept as it
       is, corrected (and exactly what changed) or written from scratch.
    2. ALTERNATE SOLUTION slides, directly after the primary solution, tagged "Alternate
       Solution" in the header - only for the questions that carry "alternate_solutions" in
       the JSON (i.e. the ones the user asked for).
    3. USEFUL RESULT / PATTERN slides after that - only for questions that carry
       "useful_results".

Mathematics: every $...$ is converted to a native PowerPoint equation (pandoc), switchable
between LaTeX and 2-D inside PowerPoint. --equations latex keeps plain $...$ text.

The JSON schema is documented in references/content_schema.md of the skill.
"""

import argparse
import json
import math
import os
import re
import sys
from copy import deepcopy

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
except ImportError:
    sys.exit("python-pptx is not installed. Run:  pip install python-pptx")

# ----------------------------------------------------------------------------
# Style - taken from the sample deck (PRT 06 Paper 01)
# ----------------------------------------------------------------------------
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x11, 0x18, 0x27)
Q_FILL = RGBColor(0xE0, 0xEC, 0xFF)    # a shade stronger than the sample so it survives projectors
Q_LINE = RGBColor(0x1D, 0x4E, 0xD8)
BOX_LINE = RGBColor(0x6B, 0x72, 0x80)    # darker grey: the sample's light grey vanishes when projected
BOX_LINE_PT = 1.5
ANS_FILL = RGBColor(0xDC, 0xFC, 0xE7)
KEY_FILL = RGBColor(0xFD, 0xE6, 0x8A)    # key-idea strip: warm amber, distinct from answer green
KEY_LINE = RGBColor(0xD9, 0x77, 0x06)
KEY_INK = RGBColor(0x78, 0x35, 0x0F)
ANS_GREEN = RGBColor(0x15, 0x80, 0x3D)

SLIDE_W, SLIDE_H = 13.333, 7.5          # inches
HEADER_H = 1.2
BOX_X, BOX_W = 0.8, 11.73
INSET_LR, INSET_T = 0.25, 0.15
BOTTOM_LIMIT = 7.3                       # keep boxes above this line

# All body text is 24 pt for classroom readability (user's instruction).
Q_FONT, Q_FACE = 24, "Arial"             # question box
INFO_FONT, SOL_FONT, TXT_FACE = 24, 24, "Arial"   # sans-serif throughout; equations get Cambria Math
ANS_FONT = 28
MIN_Q_FONT = 20                          # only used if a question cannot fit a slide at 24 pt

# Level colours, saturated enough to stay distinct on a projector (green -> red)
LEVEL_FILLS = {0: "93C47D", 1: "B6D7A8", 2: "FFD966", 3: "F6B26B", 4: "E88080", 5: "CC4125"}
LEVEL_TEXT = {0: "111827", 1: "111827", 2: "111827", 3: "111827", 4: "111827", 5: "FFFFFF"}

DEFAULT_LABELS = {
    "0": "Easier than JEE Main Level",
    "0.5": "Easier than JEE Main Level",
    "1": "JEE Main Level",
    "1.5": "A little above JEE Main Level",
    "2": "A level below Average JEE Advanced Level",
    "2.5": "A little easier than Average JEE Advanced Level",
    "3": "Average JEE Advanced Level",
    "3.5": "Average JEE Advanced Level",
    "4": "A level above Average JEE Advanced Level",
    "4.5": "A level above Average JEE Advanced Level",
    "5": "Very difficult than JEE Advanced",
}

# Marks-weighted difficulty index of official JEE Advanced maths papers, from the user's own
# ratings (see the jee-advanced-difficulty-calibrated skill). Override with "reference_indices"
# in the JSON when a new year is added.
DEFAULT_REFERENCE = [
    ["JEE Advanced 2021", 2.76],
    ["JEE Advanced 2023", 3.19],
    ["JEE Advanced 2024", 2.59],
    ["JEE Advanced 2025", 3.21],
    ["JEE Advanced 2026", 3.02],
]

# Spread groups for the summary slide (lower bound inclusive, upper exclusive)
SPREAD = [
    ("JEE Main or easier", 0, 2),
    ("Below average", 2, 3),
    ("Average Advanced", 3, 4),
    ("Above average", 4, 5),
    ("Very difficult", 5, 6),
]


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def fmt_level(x):
    x = float(x)
    return str(int(x)) if x == int(x) else f"{x:g}"


def label_for(level, labels):
    key = fmt_level(level)
    if key in labels:
        return labels[key]
    return labels.get(str(int(math.floor(float(level)))), "")


MATH = re.compile(r"\$[^$]+\$")


def visible_len(text, raw=False):
    """Rough rendered length of a string whose maths is LaTeX between $...$."""
    def squash(m):
        s = m.group(0)[1:-1]
        s = re.sub(r"\\(left|right|displaystyle|dfrac|tfrac|mathrm|operatorname|text|mathbb|mathcal)", "", s)
        s = re.sub(r"\\[a-zA-Z]+", "x", s)          # \alpha -> one glyph
        s = re.sub(r"[{}^_\\]", "", s)
        return s
    # Size for the converted deck (what the class sees), with a margin because the raw
    # LaTeX shown before conversion is longer and fractions make lines taller.
    if raw:
        return len(text)
    vis = len(MATH.sub(squash, text))
    return vis + (len(text) - vis) // 4


def n_fracs(text):
    return len(re.findall(r"\\(d|t)?frac", text))


def est_height(paragraphs, font_pt, width_in, para_gap_pt=0, raw=False):
    """Estimated height (inches) of paragraphs in a box of the given inner width."""
    cpl = max(20, int(width_in * 72 / (font_pt * 0.52)))  # characters per line
    line_h = font_pt * 1.22 / 72
    h = 0.0
    for p in paragraphs:
        lines = max(1, math.ceil(visible_len(p, raw) / cpl))
        h += lines * line_h + min(n_fracs(p), lines) * 0.12 + para_gap_pt / 72
    return h + 2 * INSET_T + 0.05


def add_box(slide, x, y, w, h, *, rounded=False, fill=WHITE, line=BOX_LINE, line_pt=BOX_LINE_PT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_pt)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(INSET_LR)
    tf.margin_top = Inches(INSET_T)
    tf.margin_bottom = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return shape


def set_text_defaults(shape, size, color):
    """Make the shape's *default* text dark (or the given colour) and the given size.

    PowerPoint autoshapes carry a <p:style> whose font colour is the theme's light colour
    (white). Runs written by this script have explicit colours, but runs created later - for
    example by a LaTeX-to-equation converter - inherit the default and turn white. Removing the
    style's font reference and setting a list-style default fixes that for any later text."""
    sp = shape._element
    style = sp.find(qn("p:style"))
    if style is not None:
        font_ref = style.find(qn("a:fontRef"))
        if font_ref is not None:
            for child in list(font_ref):
                font_ref.remove(child)
            clr = font_ref.makeelement(qn("a:srgbClr"), {"val": str(color)})
            font_ref.append(clr)
    txBody = sp.find(qn("p:txBody"))
    lst = txBody.find(qn("a:lstStyle"))
    if lst is None:
        lst = txBody.makeelement(qn("a:lstStyle"), {})
        txBody.insert(1, lst)
    for child in list(lst):
        lst.remove(child)
    lvl = lst.makeelement(qn("a:lvl1pPr"), {})
    d = lvl.makeelement(qn("a:defRPr"), {"sz": str(int(size * 100))})
    fill = d.makeelement(qn("a:solidFill"), {})
    fill.append(fill.makeelement(qn("a:srgbClr"), {"val": str(color)}))
    d.append(fill)
    lvl.append(d)
    lst.append(lvl)


def end_para_props(p, size, color):
    """Paragraph-level default run properties (as in the user's sample deck) and the end-mark
    properties, so text or equations inserted later inherit the right colour and size."""
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn("a:defRPr"))
    if old is not None:
        pPr.remove(old)
    d = pPr.makeelement(qn("a:defRPr"), {"sz": str(int(size * 100))})
    f = d.makeelement(qn("a:solidFill"), {})
    f.append(f.makeelement(qn("a:srgbClr"), {"val": str(color)}))
    d.append(f)
    # defRPr must come after spacing/bullet children of pPr
    pPr.append(d)
    epr = p._p.get_or_add_endParaRPr()
    epr.set("sz", str(int(size * 100)))
    for child in list(epr):
        epr.remove(child)
    fill = epr.makeelement(qn("a:solidFill"), {})
    fill.append(fill.makeelement(qn("a:srgbClr"), {"val": str(color)}))
    epr.append(fill)


def make_bullet(p, size):
    """Real bullet with a hanging indent, so wrapped lines align under the text."""
    pPr = p._p.get_or_add_pPr()
    indent = int(Pt(size) * 0.9)
    pPr.set("marL", str(indent))
    pPr.set("indent", str(-indent))
    defrpr = pPr.find(qn("a:defRPr"))
    for tag, attrs in (("a:buFont", {"typeface": "Arial"}), ("a:buChar", {"char": "•"})):
        el = pPr.makeelement(qn(tag), attrs)
        if defrpr is not None:
            defrpr.addprevious(el)
        else:
            pPr.append(el)


def write_paras(shape, paras, *, size, face, color=INK, bold=False, align=PP_ALIGN.LEFT,
                space_after=0, bullets=False):
    tf = shape.text_frame
    tf.clear()
    set_text_defaults(shape, size, color)
    for i, text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        add_rich(p, text, size=size, face=face, color=color, bold=bold, shape=shape)
        end_para_props(p, size, color)
        if bullets:
            make_bullet(p, size)


def level_colours(level):
    k = min(5, int(math.floor(float(level))))
    return RGBColor.from_string(LEVEL_FILLS[k]), RGBColor.from_string(LEVEL_TEXT[k])


def header(slide, title, q=None, tag=None, tag_kind="amber"):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(HEADER_H))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    bar.shadow.inherit = False
    tf = bar.text_frame
    tf.margin_left = Inches(0.8)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    write_paras(bar, [title], size=32, face="Arial", color=WHITE, bold=True, align=PP_ALIGN.LEFT)
    if q is not None:
        # Badge on every slide of a question: its level (in the level colour) and marks
        w, h = 4.2, 0.66
        fill, ink = level_colours(q["level"])
        badge = add_box(slide, SLIDE_W - 0.5 - w, (HEADER_H - h) / 2, w, h, rounded=True,
                        fill=fill, line=WHITE, line_pt=1.5)
        badge.adjustments[0] = 0.5
        badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        badge.text_frame.margin_top = badge.text_frame.margin_bottom = Inches(0.02)
        write_paras(badge, [f"Level {fmt_level(q['level'])}  ·  {q['marks']} Marks"],
                    size=24, face="Arial", color=ink, bold=True, align=PP_ALIGN.CENTER)
    if tag:
        # Tag pill between the title and the level badge: "Alternate Solution" (amber) or
        # "Useful Result / Pattern" (blue)
        w, h = 4.6, 0.62
        fill, line, ink = (KEY_FILL, KEY_LINE, KEY_INK) if tag_kind == "amber" else (Q_FILL, Q_LINE, NAVY)
        pill = add_box(slide, 3.85, (HEADER_H - h) / 2, w, h, rounded=True, fill=fill, line=line, line_pt=2.0)
        pill.adjustments[0] = 0.5
        pill.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        pill.text_frame.margin_top = pill.text_frame.margin_bottom = Inches(0.02)
        write_paras(pill, [tag], size=22, face="Arial", color=ink, bold=True, align=PP_ALIGN.CENTER)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])



# ----------------------------------------------------------------------------
# Native PowerPoint equations (default)
# ----------------------------------------------------------------------------
# LaTeX between $...$ is converted to Office Math (OMML) with pandoc and embedded as real
# PowerPoint equations, exactly the way PowerPoint stores them itself. In PowerPoint each one
# can be switched between LaTeX and 2-D form: click the equation -> Equation tab ->
# choose "LaTeX" and "Linear" to see/edit LaTeX, then "Professional" to go back to 2-D.
# With --equations latex the $...$ text is left as it is (for an external converter).

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
EQ = {"mode": "latex", "omml": {}, "shapes": [], "failed": set(), "wrappers": [], "why": {}}

_ITALIC = {c: chr(0x1D44E + i) for i, c in enumerate("abcdefghijklmnopqrstuvwxyz")}
_ITALIC["h"] = "\u210E"
_ITALIC.update({c: chr(0x1D434 + i) for i, c in enumerate("ABCDEFGHIJKLMNOPQRSTUVWXYZ")})


def find_pandoc():
    try:
        import pypandoc
        path = pypandoc.get_pandoc_path()
        if path:
            return path
    except Exception:
        pass
    import shutil
    return shutil.which("pandoc")


def all_latex(data):
    found = []

    def walk(v):
        if isinstance(v, str):
            found.extend(m.group(0)[1:-1] for m in MATH.finditer(v))
        elif isinstance(v, list):
            for x in v:
                walk(x)
        elif isinstance(v, dict):
            for x in v.values():
                walk(x)
    walk(data.get("questions", []))
    return list(dict.fromkeys(found))


def convert_latex(snippets, pandoc):
    """Convert every LaTeX snippet to OMML in one pandoc run (LaTeX -> .docx -> OMML)."""
    import subprocess
    import tempfile
    import zipfile
    from lxml import etree
    tmp = tempfile.mkdtemp()
    md, docx = os.path.join(tmp, "eq.md"), os.path.join(tmp, "eq.docx")
    with open(md, "w", encoding="utf-8") as fh:
        for i, snip in enumerate(snippets):
            fh.write(f"EQ{i}\n\n${snip}$\n\n")
    subprocess.run([pandoc, md, "-f", "markdown", "-o", docx], check=True, capture_output=True)
    xml = zipfile.ZipFile(docx).read("word/document.xml")
    body = etree.fromstring(xml)
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    result, current = {}, None
    for p in body.iter(f"{{{W}}}p"):
        text = "".join(t.text or "" for t in p.iter(f"{{{W}}}t"))
        maths = p.findall(f"{{{NS_M}}}oMath") + p.findall(f".//{{{NS_M}}}oMath")
        if text.startswith("EQ") and text[2:].isdigit() and not maths:
            current = int(text[2:])
        elif current is not None and maths:
            result[snippets[current]] = maths[0]
            current = None
    return result


def math_rpr(size, color, italic):
    from lxml import etree
    r = etree.Element(f"{{{NS_A}}}rPr", {"lang": "en-IN", "sz": str(int(size * 100))})
    if italic:
        r.set("i", "1")
    fill = etree.SubElement(r, f"{{{NS_A}}}solidFill")
    etree.SubElement(fill, f"{{{NS_A}}}srgbClr", {"val": str(color)})
    etree.SubElement(r, f"{{{NS_A}}}latin", {"typeface": "Cambria Math", "panose": "02040503050406030204",
                                             "pitchFamily": "18", "charset": "0"})
    return r


# ----------------------------------------------------------------------------
# OMML structure rules - generated from the OMML schema (shared-math.xsd, ISO/IEC 29500)
# ----------------------------------------------------------------------------
# PowerPoint is strict about the ORDER of child elements inside equations; Word is not.
# Pandoc writes some elements in an order Word accepts but PowerPoint rejects - the known case
# is every matrix / determinant / array / cases environment: <m:mcPr><m:mcJc/><m:count/>
# (the schema requires <m:count/> first). PowerPoint answers with "repair" and empties the slide.
# normalise_omml() re-orders the children of every element whose order the schema fixes, and
# only the elements listed in OMML_CTRLPR_OK may carry the <m:ctrlPr> run-formatting element.
OMML_ORDER = {'acc': ['accPr', 'e'],
 'accPr': ['chr', 'ctrlPr'],
 'bar': ['barPr', 'e'],
 'barPr': ['pos', 'ctrlPr'],
 'borderBox': ['borderBoxPr', 'e'],
 'borderBoxPr': ['hideTop', 'hideBot', 'hideLeft', 'hideRight', 'strikeH', 'strikeV', 'strikeBLTR',
                 'strikeTLBR', 'ctrlPr'],
 'box': ['boxPr', 'e'],
 'boxPr': ['opEmu', 'noBreak', 'diff', 'brk', 'aln', 'ctrlPr'],
 'd': ['dPr', 'e'],
 'dPr': ['begChr', 'sepChr', 'endChr', 'grow', 'shp', 'ctrlPr'],
 'eqArr': ['eqArrPr', 'e'],
 'eqArrPr': ['baseJc', 'maxDist', 'objDist', 'rSpRule', 'rSp', 'ctrlPr'],
 'f': ['fPr', 'num', 'den'],
 'fPr': ['type', 'ctrlPr'],
 'func': ['funcPr', 'fName', 'e'],
 'groupChr': ['groupChrPr', 'e'],
 'groupChrPr': ['chr', 'pos', 'vertJc', 'ctrlPr'],
 'limLow': ['limLowPr', 'e', 'lim'],
 'limUpp': ['limUppPr', 'e', 'lim'],
 'm': ['mPr', 'mr'],
 'mPr': ['baseJc', 'plcHide', 'rSpRule', 'cGpRule', 'rSp', 'cSp', 'cGp', 'mcs', 'ctrlPr'],
 'mcPr': ['count', 'mcJc'],
 'nary': ['naryPr', 'sub', 'sup', 'e'],
 'naryPr': ['chr', 'limLoc', 'grow', 'subHide', 'supHide', 'ctrlPr'],
 'oMathPara': ['oMathParaPr', 'oMath'],
 'phant': ['phantPr', 'e'],
 'phantPr': ['show', 'zeroWid', 'zeroAsc', 'zeroDesc', 'transp', 'ctrlPr'],
 'rad': ['radPr', 'deg', 'e'],
 'radPr': ['degHide', 'ctrlPr'],
 'sPre': ['sPrePr', 'sub', 'sup', 'e'],
 'sSub': ['sSubPr', 'e', 'sub'],
 'sSubSup': ['sSubSupPr', 'e', 'sub', 'sup'],
 'sSubSupPr': ['alnScr', 'ctrlPr'],
 'sSup': ['sSupPr', 'e', 'sup']}

OMML_CTRLPR_OK = set(['accPr', 'barPr', 'borderBoxPr', 'boxPr', 'dPr', 'eqArrPr', 'fPr', 'funcPr', 'groupChrPr', 'limLowPr',
 'limUppPr', 'mPr', 'naryPr', 'phantPr', 'radPr', 'sPrePr', 'sSubPr', 'sSubSupPr', 'sSupPr'])


def normalise_omml(om):
    """Put the children of every OMML element into the order the schema requires."""
    from lxml import etree
    mns = "{" + NS_M + "}"
    RPR_ORDER = ["lit", "nor", "scr", "sty", "brk", "aln"]     # m:rPr: lit, (nor | scr+sty), brk, aln
    for el in om.iter():
        if not isinstance(el.tag, str) or not el.tag.startswith(mns):
            continue
        if el.tag == mns + "rPr":
            # \text{...} gives <m:nor/> AND <m:sty/>; the schema makes them mutually exclusive
            # ("normal text" is already upright), so drop the style when nor is present.
            kids = [c for c in el if isinstance(c.tag, str) and c.tag.startswith(mns)]
            names = [c.tag[len(mns):] for c in kids]
            if "nor" in names:
                for c in kids:
                    if c.tag[len(mns):] in ("sty", "scr"):
                        el.remove(c)
                kids = [c for c in kids if c.tag[len(mns):] not in ("sty", "scr")]
            if all(c.tag[len(mns):] in RPR_ORDER for c in kids):
                ordered = sorted(kids, key=lambda c: RPR_ORDER.index(c.tag[len(mns):]))
                if ordered != kids:
                    for c in kids:
                        el.remove(c)
                    for c in ordered:
                        el.append(c)
            continue
        order = OMML_ORDER.get(el.tag[len(mns):])
        if not order:
            continue
        kids = [c for c in el if isinstance(c.tag, str)]
        if not all(c.tag.startswith(mns) and c.tag[len(mns):] in order for c in kids):
            continue                      # unknown content: leave untouched rather than guess
        ordered = sorted(kids, key=lambda c: order.index(c.tag[len(mns):]))   # stable sort
        if ordered != kids:
            for c in kids:
                el.remove(c)
            for c in ordered:
                el.append(c)
    return om


OMML_KNOWN = set(['acc', 'accPr', 'aln', 'alnScr', 'argPr', 'argSz', 'bar', 'barPr', 'baseJc', 'begChr', 'borderBox',
 'borderBoxPr', 'box', 'boxPr', 'brk', 'brkBin', 'brkBinSub', 'cGp', 'cGpRule', 'cSp', 'chr', 'count',
 'ctrlPr', 'd', 'dPr', 'defJc', 'deg', 'degHide', 'den', 'diff', 'dispDef', 'e', 'endChr', 'eqArr', 'eqArrPr',
 'f', 'fName', 'fPr', 'func', 'funcPr', 'groupChr', 'groupChrPr', 'grow', 'hideBot', 'hideLeft', 'hideRight',
 'hideTop', 'intLim', 'interSp', 'intraSp', 'jc', 'lMargin', 'lim', 'limLoc', 'limLow', 'limLowPr', 'limUpp',
 'limUppPr', 'lit', 'm', 'mPr', 'mathFont', 'mathPr', 'maxDist', 'mc', 'mcJc', 'mcPr', 'mcs', 'mr', 'nary',
 'naryLim', 'naryPr', 'noBreak', 'nor', 'num', 'oMath', 'oMathPara', 'oMathParaPr', 'objDist', 'opEmu',
 'phant', 'phantPr', 'plcHide', 'pos', 'postSp', 'preSp', 'r', 'rMargin', 'rPr', 'rSp', 'rSpRule', 'rad',
 'radPr', 'sPre', 'sPrePr', 'sSub', 'sSubPr', 'sSubSup', 'sSubSupPr', 'sSup', 'sSupPr', 'scr', 'sepChr',
 'show', 'shp', 'smallFrac', 'strikeBLTR', 'strikeH', 'strikeTLBR', 'strikeV', 'sty', 'sub', 'subHide', 'sup',
 'supHide', 't', 'transp', 'type', 'vertJc', 'wrapIndent', 'wrapRight', 'zeroAsc', 'zeroDesc', 'zeroWid'])


def omml_problem(om):
    """Return a short description of what is structurally wrong with an equation, or None if it
    follows the schema rules this script knows (known elements only, children in schema order,
    <m:ctrlPr> only where allowed). Equations that fail are NOT embedded: the LaTeX text stays on
    the slide and a warning is printed, so one bad equation can never blank a slide in PowerPoint."""
    mns = "{" + NS_M + "}"
    for el in om.iter():
        if not isinstance(el.tag, str) or not el.tag.startswith(mns):
            continue
        name = el.tag[len(mns):]
        if name not in OMML_KNOWN:
            return f"unknown element m:{name}"
        kids = [c for c in el if isinstance(c.tag, str) and c.tag.startswith(mns)]
        kid_names = [c.tag[len(mns):] for c in kids]
        if "ctrlPr" in kid_names and name not in OMML_CTRLPR_OK:
            return f"m:ctrlPr is not allowed inside m:{name}"
        if name == "rPr":
            if "nor" in kid_names and ("sty" in kid_names or "scr" in kid_names):
                return "m:rPr has both nor and sty/scr"
            continue
        order = OMML_ORDER.get(name)
        if order:
            if any(k not in order for k in kid_names):
                return f"m:{name} has a child the schema does not allow"
            idx = [order.index(k) for k in kid_names]
            if idx != sorted(idx):
                return f"children of m:{name} are out of schema order"
    return None


def keep_safe(omml):
    """Normalise every converted equation and drop those that still look unsafe."""
    ok = {}
    for latex, om in omml.items():
        om = normalise_omml(om)
        why = omml_problem(om)
        if why:
            EQ["why"][latex] = why
        else:
            ok[latex] = om
    return ok


def styled_omml(latex, size, color):
    """Copy of the converted equation with PowerPoint run formatting (colour, size, font)."""
    from lxml import etree
    om = normalise_omml(deepcopy(EQ["omml"][latex]))
    for r in om.iter(f"{{{NS_M}}}r"):
        rpr = r.find(f"{{{NS_M}}}rPr")
        sty = rpr.find(f"{{{NS_M}}}sty") if rpr is not None else None
        italic = sty is None or sty.get(f"{{{NS_M}}}val") not in ("p", "b")
        t = r.find(f"{{{NS_M}}}t")
        if italic and t is not None and t.text:
            t.text = "".join(_ITALIC.get(c, c) for c in t.text)
        pos = 1 if rpr is not None else 0
        r.insert(pos, math_rpr(size, color, italic))
    for el in list(om.iter()):
        if not isinstance(el.tag, str) or not el.tag.startswith(f"{{{NS_M}}}"):
            continue
        # <m:ctrlPr> (run formatting of the structure) is legal only in the elements the schema
        # lists, and always as the LAST child; elsewhere (e.g. <m:mcPr>) it makes PowerPoint repair.
        if etree.QName(el).localname in OMML_CTRLPR_OK and el.find(f"{{{NS_M}}}ctrlPr") is None:
            ctrl = etree.SubElement(el, f"{{{NS_M}}}ctrlPr")
            ctrl.append(math_rpr(size, color, True))
    return om


def add_rich(p, text, *, size, face, color, bold=False, shape=None):
    """Add text to paragraph p; $...$ becomes a native equation in native mode."""
    from lxml import etree
    parts = [text] if EQ["mode"] != "native" else re.split(r"(\$[^$]+\$)", text)
    for part in parts:
        if not part:
            continue
        latex = part[1:-1] if (part.startswith("$") and part.endswith("$") and len(part) > 1) else None
        if latex is not None and latex in EQ["omml"]:
            wrapper = etree.SubElement(p._p, f"{{{NS_A14}}}m", nsmap={"a14": NS_A14})
            wrapper.append(styled_omml(latex, size, color))
            EQ["wrappers"].append((wrapper, latex, size, color))
            if shape is not None and shape not in EQ["shapes"]:
                EQ["shapes"].append(shape)
            continue
        if latex is not None:
            EQ["failed"].add(latex)
        r = p.add_run()
        r.text = part
        f = r.font
        f.size, f.name, f.bold = Pt(size), face, bold
        f.color.rgb = color


def wrap_math_shapes():
    """PowerPoint keeps text boxes that contain equations inside mc:AlternateContent.
    Choice = the box with native equations (used by PowerPoint 2010 and later).
    Fallback = the same box with the LaTeX as plain text, for apps that cannot show
    Office equations (LibreOffice, some previewers); PowerPoint replaces it on save."""
    from lxml import etree
    for shape in EQ["shapes"]:
        sp = shape._element
        fallback_sp = deepcopy(sp)
        originals = list(sp.iter(f"{{{NS_A14}}}m"))
        copies = list(fallback_sp.iter(f"{{{NS_A14}}}m"))
        for orig, cp in zip(originals, copies):
            info = next((w for w in EQ["wrappers"] if w[0] is orig), None)
            latex, size, color = (info[1], info[2], info[3]) if info else ("", 24, INK)
            r = etree.Element(f"{{{NS_A}}}r")
            r.append(math_rpr(size, color, False))
            etree.SubElement(r, f"{{{NS_A}}}t").text = f"${latex}$"
            cp.getparent().replace(cp, r)
        parent = sp.getparent()
        idx = parent.index(sp)
        ac = etree.Element(f"{{{NS_MC}}}AlternateContent", nsmap={"mc": NS_MC, "a14": NS_A14})
        choice = etree.SubElement(ac, f"{{{NS_MC}}}Choice", {"Requires": "a14"})
        fb = etree.SubElement(ac, f"{{{NS_MC}}}Fallback")
        parent.remove(sp)
        choice.append(sp)
        fb.append(fallback_sp)
        parent.insert(idx, ac)


# ----------------------------------------------------------------------------
# Slide builders
# ----------------------------------------------------------------------------
def title_slide(prs, title, subtitle):
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    bg.shadow.inherit = False
    tf = bg.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    write_paras(bg, [title], size=40, face="Arial", color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    if subtitle:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = subtitle
        r.font.size, r.font.name, r.font.color.rgb = Pt(22), "Arial", WHITE


def build_sections(data, qs):
    """Sections from the JSON ("sections"), or consecutive runs of the same type and marks."""
    if data.get("sections"):
        out = []
        for sec in data["sections"]:
            out.append({"title": sec["title"], "from": int(sec["from"]), "to": int(sec["to"]),
                        "marking": sec.get("marking") or []})
        return out
    out = []
    for q in qs:
        key = (q["type_label"].strip(), q["marks"])
        if out and out[-1]["key"] == key:
            out[-1]["to"] = int(q["q_no"])
        else:
            out.append({"key": key, "title": q["type_label"].strip(), "from": int(q["q_no"]),
                        "to": int(q["q_no"]), "marking": [f"{q['marks']} marks each"]})
    return out


def section_slide(prs, number, sec):
    """Divider before each section: name, question range and marking scheme."""
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    bg.shadow.inherit = False
    rng = f"Q{sec['from']:02d}" if sec["from"] == sec["to"] else f"Q{sec['from']:02d} – Q{sec['to']:02d}"
    box = s.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(SLIDE_W - 2.0), Inches(4.5))
    box.text_frame.word_wrap = True
    lines = [f"Section {number}", sec["title"], rng] + list(sec["marking"])
    write_paras(box, lines, size=28, face="Arial", color=WHITE, align=PP_ALIGN.CENTER, space_after=10)
    ps = box.text_frame.paragraphs
    ps[0].runs[0].font.size = Pt(28)
    ps[0].runs[0].font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
    ps[1].runs[0].font.size, ps[1].runs[0].font.bold = Pt(44), True
    ps[2].runs[0].font.size, ps[2].runs[0].font.bold = Pt(32), True


def question_text(q):
    paras = []
    if q.get("paragraph"):
        paras += [t for t in q["paragraph"].split("\n") if t.strip()]
    paras += [t for t in q["question"].split("\n") if t.strip()]
    opts = q.get("options") or []
    if opts:
        letters = "ABCDEFGH"
        tagged = [f"({letters[i]}) {o}" for i, o in enumerate(opts)]
        one_line = "   ".join(tagged)
        if visible_len(one_line) <= 80:
            paras.append(one_line)
        else:
            paras += tagged
    return paras


def info_lines(q, labels):
    return [f"{q['type_label']}({q['marks']} Marks)",
            f"Concept(s) : {', '.join(q['concepts'])}",
            f"Difficulty Level: Level {fmt_level(q['level'])}({label_for(q['level'], labels)})"]


def info_height(info):
    return max(1.5, min(2.4, est_height(info, INFO_FONT, BOX_W - 2 * INSET_LR, para_gap_pt=2) + 0.1))


def question_slide(prs, q, labels, warnings):
    """Question box plus info box at 24 pt. Returns True if the info box had to move to the
    first solution slide because the question is too long to share its slide."""
    s = blank(prs)
    header(s, f"Question {int(q['q_no']):02d}", q)
    paras = question_text(q)
    inner_w = BOX_W - 2 * INSET_LR
    info = info_lines(q, labels)
    info_h, gap, top = info_height(info), 0.25, 1.4

    def q_height(font, raw):
        return max(1.9, est_height(paras, font, inner_w, para_gap_pt=5, raw=raw) + 0.15)

    # 1) question + info together at 24 pt (sized for raw LaTeX, then for converted maths)
    # Same sizing in both equation modes, so boxes sit exactly where they did in the
    # LaTeX-text deck (the user's reference layout).
    for raw in (True, False):
        q_h = q_height(Q_FONT, raw)
        if top + q_h + gap + info_h <= BOTTOM_LIMIT:
            _draw_question(s, paras, q_h, Q_FONT)
            ib = add_box(s, BOX_X, top + q_h + gap, BOX_W, info_h)
            write_paras(ib, info, size=INFO_FONT, face=TXT_FACE, space_after=2)
            return False
    # 2) question alone on the slide at 24 pt; info box moves to the first solution slide
    for font in (Q_FONT, 22, MIN_Q_FONT):
        q_h = q_height(font, False)
        if top + q_h <= BOTTOM_LIMIT or font == MIN_Q_FONT:
            break
    if font < Q_FONT:
        warnings.append(f"Q{q['q_no']}: very long question - font reduced to {font} pt.")
    _draw_question(s, paras, min(q_h, BOTTOM_LIMIT - top), font)
    return True


def _draw_question(s, paras, q_h, font):
    qb = add_box(s, BOX_X, 1.4, BOX_W, q_h, rounded=True, fill=Q_FILL, line=Q_LINE, line_pt=2.0)
    write_paras(qb, paras, size=font, face=Q_FACE, space_after=5)


ANS_Y, ANS_H = 6.45, 0.85
KEY_LABEL = "Key idea:  "
IDEA_LABEL = "Idea:  "


def chunk_solution(bullets, max_bullets, first_top=1.48, has_final=True):
    """Split bullets into slides: <= max_bullets each and within the box height. The first
    slide's box starts at first_top (lower when an info box / strip sits above it) and, when
    has_final, the last slide leaves room for the green answer box."""
    inner_w = BOX_W - 2 * INSET_LR
    height = lambda c: est_height(["•  " + t for t in c], SOL_FONT, inner_w, 6)
    room = lambda top, last: ((ANS_Y - 0.12) if (last and has_final) else BOTTOM_LIMIT) - top
    chunks, cur = [], []
    for b in bullets:
        top = first_top if not chunks else 1.48
        trial = cur + [b]
        if cur and (len(trial) > max_bullets or height(trial) > room(top, False)):
            chunks.append(cur)
            cur = [b]
        else:
            cur = trial
    if cur:
        chunks.append(cur)
    if not chunks:
        return [[]]
    # the final chunk must also fit above the answer box; if not, add a slide
    last_top = first_top if len(chunks) == 1 else 1.48
    if height(chunks[-1]) > room(last_top, True) and len(chunks[-1]) > 1:
        last = chunks.pop()
        half = (len(last) + 1) // 2
        chunks += [last[:half], last[half:]]
    # spread bullets evenly over the slides (avoids a lone bullet on the last slide)
    n = len(chunks)
    if n > 1:
        base, extra = divmod(len(bullets), n)
        even, i = [], 0
        for k in range(n):
            size_k = base + (1 if k < extra else 0)
            even.append(bullets[i:i + size_k])
            i += size_k
        ok = all(len(c) <= max_bullets and
                 height(c) <= room(first_top if k == 0 else 1.48, k == n - 1)
                 for k, c in enumerate(even))
        if ok:
            chunks = even
    return chunks


def key_strip_height(text, label=KEY_LABEL):
    return max(0.75, est_height([label + text], SOL_FONT, BOX_W - 2 * INSET_LR - 0.1) + 0.05)


def key_strip(slide, text, y, h, label=KEY_LABEL):
    """Amber strip at the top of the first slide of a solution: the decisive step ("Key idea")
    or, on an alternate solution, the idea of the alternate method ("Idea")."""
    box = add_box(slide, BOX_X, y, BOX_W, h, rounded=True, fill=KEY_FILL, line=KEY_LINE, line_pt=2.5)
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    box.text_frame.margin_top = Inches(0.06)
    write_paras(box, [label], size=SOL_FONT, face=TXT_FACE, color=KEY_INK, bold=True)
    p = box.text_frame.paragraphs[0]
    # Take the end-of-paragraph mark out first: python-pptx inserts new text runs before it
    # while equations are appended after the last child, so leaving it in place would put the
    # strip's text and equations out of order.
    epr = p._p.find(qn("a:endParaRPr"))
    if epr is not None:
        p._p.remove(epr)
    add_rich(p, text, size=SOL_FONT, face=TXT_FACE, color=KEY_INK, shape=box)
    if epr is not None:
        p._p.append(epr)


def solution_slides(prs, q, labels, max_bullets, warnings, info_first=False, *, bullets=None,
                    tag=None, tag_kind="amber", strip=None, strip_label=KEY_LABEL,
                    final_label=None, notes=None, strip_required=False, name="solution"):
    """Slides of bullets for one question. Used three ways:
       primary solution  - key-idea strip, green "Correct Answer" box, speaker notes
       alternate solution - tag pill, optional "Idea" strip, green "Same answer" box
       useful results    - tag pill (blue), bullets only"""
    bullets = bullets or []
    qn = int(q["q_no"])
    if not bullets:
        warnings.append(f"Q{qn}: no {name} bullets in the JSON.")
    info = info_lines(q, labels)
    info_h = info_height(info)
    key = (strip or "").strip()
    if strip_required and not key:
        warnings.append(f"Q{qn}: no key_idea in the JSON - strip skipped.")
    key_y = 1.4 + info_h + 0.2 if info_first else 1.4
    key_h = key_strip_height(key, strip_label) if key else 0
    first_top = key_y + key_h + 0.15 if key else (1.4 + info_h + 0.2 if info_first else 1.48)
    has_final = final_label is not None
    chunks = chunk_solution(bullets, max_bullets, first_top, has_final)
    for i, chunk in enumerate(chunks):
        s = blank(prs)
        header(s, f"Question {qn:02d}", q, tag=tag, tag_kind=tag_kind)
        if i == 0 and notes:
            s.notes_slide.notes_text_frame.text = notes
        top = 1.48
        if i == 0 and info_first:
            ib = add_box(s, BOX_X, 1.4, BOX_W, info_h)
            write_paras(ib, info, size=INFO_FONT, face=TXT_FACE, space_after=2)
        if i == 0 and key:
            key_strip(s, key, key_y, key_h, strip_label)
        if i == 0:
            top = first_top
        last = i == len(chunks) - 1
        bottom = (ANS_Y - 0.12) if (last and has_final) else BOTTOM_LIMIT
        box = add_box(s, BOX_X, top, BOX_W, bottom - top)
        write_paras(box, chunk, size=SOL_FONT, face=TXT_FACE, space_after=6, bullets=True)
        if last and has_final:
            ans = add_box(s, BOX_X, ANS_Y, BOX_W, ANS_H, rounded=True, fill=ANS_FILL, line=ANS_GREEN, line_pt=2.0)
            ans.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            ans.text_frame.margin_top = Inches(0.02)
            write_paras(ans, [f"{final_label}: {q.get('answer', '')}"], size=ANS_FONT, face=TXT_FACE,
                        color=ANS_GREEN, bold=True, align=PP_ALIGN.CENTER)
            if not q.get("answer"):
                warnings.append(f"Q{qn}: no answer in the JSON.")


SOURCE_TEXT = {
    "pdf": "Solution copied from the PDF in full and checked step by step - it is correct.",
    "corrected": "Solution taken from the PDF but CORRECTED by Claude.",
    "written": "The PDF has no usable solution for this question; the solution was written by Claude.",
}


def speaker_note(q, warnings):
    src = (q.get("solution_source") or "").strip().lower()
    if src not in SOURCE_TEXT:
        warnings.append(f"Q{q['q_no']}: solution_source must be pdf / corrected / written (got {src or 'nothing'}).")
        return None
    note = SOURCE_TEXT[src]
    cn = (q.get("correction_note") or "").strip()
    if src == "corrected":
        if not cn:
            warnings.append(f"Q{q['q_no']}: solution_source is 'corrected' but correction_note is empty.")
        note += "\nWHAT WAS WRONG IN THE PDF AND WHAT CHANGED: " + cn
    return note


def question_solution_slides(prs, q, labels, max_bullets, warnings, info_first):
    """Primary solution, then each alternate solution, then useful results - in that order."""
    if len(q.get("solution") or []) < 4:
        warnings.append(f"Q{q['q_no']}: only {len(q.get('solution') or [])} solution bullets - a complete "
                        f"solution normally has more (every step of the PDF's working).")
    solution_slides(prs, q, labels, max_bullets, warnings, info_first=info_first, bullets=q.get("solution"),
                    strip=q.get("key_idea"), strip_label=KEY_LABEL, final_label="Correct Answer",
                    notes=speaker_note(q, warnings), strip_required=True)
    alts = q.get("alternate_solutions") or []
    for k, alt in enumerate(alts, 1):
        if isinstance(alt, list):
            alt = {"solution": alt}
        tag = "Alternate Solution" + (f" {k}" if len(alts) > 1 else "")
        solution_slides(prs, q, labels, max_bullets, warnings, bullets=alt.get("solution"), tag=tag,
                        strip=alt.get("idea"), strip_label=IDEA_LABEL, final_label="Same answer",
                        name="alternate solution")
    if q.get("useful_results"):
        solution_slides(prs, q, labels, max_bullets, warnings, bullets=q["useful_results"],
                        tag="Useful Result / Pattern", tag_kind="blue", name="useful result")


def table_slides(prs, qs, labels, rows_per_slide=5):
    """Native, editable table sized for a large classroom: up to 5 questions per slide, 24 pt,
    questions spread evenly across the table slides."""
    cols = [("Q No", 1.1), ("Question Type", 3.2), ("Marks", 1.4), ("Level", 1.3),
            ("Level Description", 5.33)]
    n_parts = math.ceil(len(qs) / rows_per_slide)
    base, extra = divmod(len(qs), n_parts)
    parts, start = [], 0
    for k in range(n_parts):
        size_k = base + (1 if k < extra else 0)
        parts.append(qs[start:start + size_k])
        start += size_k
    for pi, part in enumerate(parts):
        s = blank(prs)
        suffix = f" ({pi + 1}/{len(parts)})" if len(parts) > 1 else ""
        header(s, "Overall Difficulty Analysis" + suffix)
        n = len(part) + 1
        head_h, row_h = 0.75, 0.95
        gt = s.shapes.add_table(n, len(cols), Inches(0.5), Inches(1.45),
                                Inches(sum(w for _, w in cols)), Inches(head_h + row_h * (n - 1)))
        t = gt.table
        for c, (name, w) in enumerate(cols):
            t.columns[c].width = Inches(w)
        t.rows[0].height = Inches(head_h)
        for r in range(1, n):
            t.rows[r].height = Inches(row_h)

        def put(r, c, text, *, bold=False, fill=None, color=INK, align=PP_ALIGN.CENTER):
            cell = t.cell(r, c)
            cell.text = ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.1)
            p = cell.text_frame.paragraphs[0]
            p.alignment = align
            run = p.add_run()
            run.text = str(text)
            run.font.size, run.font.name, run.font.bold = Pt(24), "Arial", bold
            run.font.color.rgb = color
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill if fill is not None else WHITE

        for c, (name, _) in enumerate(cols):
            put(0, c, name, bold=True, fill=NAVY, color=WHITE)
        for r, q in enumerate(part, 1):
            lv = float(q["level"])
            lvl_fill, lvl_ink = level_colours(lv)
            put(r, 0, f"{int(q['q_no']):02d}", bold=True)
            put(r, 1, q.get("short_type") or q["type_label"], align=PP_ALIGN.LEFT)
            put(r, 2, q["marks"])
            put(r, 3, fmt_level(lv), bold=True, fill=lvl_fill, color=lvl_ink)
            put(r, 4, label_for(lv, labels), fill=lvl_fill, color=lvl_ink, align=PP_ALIGN.LEFT)


def summary_slide(prs, qs):
    s = blank(prs)
    header(s, "Overall Difficulty Analysis")
    marks = sum(float(q["marks"]) for q in qs)
    index = sum(float(q["marks"]) * float(q["level"]) for q in qs) / marks if marks else 0

    big = add_box(s, 0.8, 1.55, 5.2, 3.6, rounded=True, fill=Q_FILL, line=Q_LINE, line_pt=2.0)
    big.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    write_paras(big, ["Difficulty Index", f"{index:.2f}", "out of 5"], size=30, face="Arial",
                color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    big.text_frame.paragraphs[1].runs[0].font.size = Pt(88)
    big.text_frame.paragraphs[2].runs[0].font.size = Pt(24)
    big.text_frame.paragraphs[2].runs[0].font.bold = False

    y = 1.55
    for name, lo, hi in SPREAD:
        sel = [q for q in qs if lo <= float(q["level"]) < hi]
        fill, ink = level_colours(lo)
        card = add_box(s, 6.4, y, 6.13, 0.66, fill=fill, line=BOX_LINE)
        card.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        card.text_frame.margin_top = Inches(0.02)
        qtxt = "question" if len(sel) == 1 else "questions"
        write_paras(card, [f"{name}:  {len(sel)} {qtxt}"], size=24, face="Arial", color=ink, bold=True)
        y += 0.74
    note = add_box(s, 0.8, 5.5, 11.73, 1.2, fill=WHITE, line=WHITE)
    note.line.fill.background()
    write_paras(note, [f"Total: {len(qs)} questions, {marks:g} marks.  "
                       "Index = Σ(marks × level) ÷ Σ marks, on the 0-5 scale."],
                size=24, face=TXT_FACE, align=PP_ALIGN.CENTER)


def profile_slide(prs, qs):
    """Native column chart: one bar per question, height = level, coloured by level."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

    s = blank(prs)
    header(s, "Difficulty Profile of the Paper")
    cd = CategoryChartData()
    cd.categories = [f"Q{int(q['q_no'])}" for q in qs]
    cd.add_series("Level", [float(q["level"]) for q in qs])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.4),
                            Inches(12.33), Inches(5.85), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 40
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = "0.0#"
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold, dl.font.name = Pt(20), True, "Arial"
    dl.font.color.rgb = INK
    ser = plot.series[0]
    for i, q in enumerate(qs):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = level_colours(q["level"])[0]
        pt.format.line.color.rgb = INK
        pt.format.line.width = Pt(0.75)
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 5.5, 1
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xD1, 0xD5, 0xDB)
    va.tick_labels.font.size, va.tick_labels.font.name = Pt(20), "Arial"
    va.has_title = True
    va.axis_title.text_frame.text = "Level"
    va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    ca = ch.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.bold, ca.tick_labels.font.name = Pt(20), True, "Arial"


def comparison_slide(prs, qs, reference, paper_name):
    """Last slide: this paper's index next to the official JEE Advanced papers."""
    s = blank(prs)
    header(s, "Comparison with JEE Advanced Papers")
    marks = sum(float(q["marks"]) for q in qs)
    index = sum(float(q["marks"]) * float(q["level"]) for q in qs) / marks if marks else 0
    rows = [(name, float(v)) for name, v in reference] + [(paper_name, index)]
    n = len(rows) + 1
    col_w = (6.3, 3.4)
    head_h = 0.75
    row_h = min(0.7, (BOTTOM_LIMIT - 2.2 - head_h) / (n - 1))
    x = (SLIDE_W - sum(col_w)) / 2
    gt = s.shapes.add_table(n, 2, Inches(x), Inches(1.5), Inches(sum(col_w)),
                            Inches(head_h + row_h * (n - 1)))
    t = gt.table
    for c, w in enumerate(col_w):
        t.columns[c].width = Inches(w)
    t.rows[0].height = Inches(head_h)
    for r in range(1, n):
        t.rows[r].height = Inches(row_h)

    def put(r, c, text, *, bold=False, fill=WHITE, color=INK, align=PP_ALIGN.CENTER):
        cell = t.cell(r, c)
        cell.text = ""
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.15)
        p = cell.text_frame.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size, run.font.name, run.font.bold = Pt(26), "Arial", bold
        run.font.color.rgb = color
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill

    put(0, 0, "Paper", bold=True, fill=NAVY, color=WHITE)
    put(0, 1, "Difficulty Index", bold=True, fill=NAVY, color=WHITE)
    for r, (name, v) in enumerate(rows, 1):
        mine = r == len(rows)
        fill = Q_FILL if mine else WHITE
        put(r, 0, name, bold=mine, fill=fill, color=NAVY if mine else INK, align=PP_ALIGN.LEFT)
        put(r, 1, f"{v:.2f}", bold=True, fill=fill, color=NAVY if mine else INK)
    note_y = 1.5 + head_h + row_h * (n - 1) + 0.2
    note = s.shapes.add_textbox(Inches(0.8), Inches(note_y), Inches(11.73), Inches(0.5))
    write_paras(note, ["Marks-weighted difficulty index on the 0-5 scale"], size=20,
                face=TXT_FACE, align=PP_ALIGN.CENTER)


def air_slide(prs, air):
    s = blank(prs)
    header(s, "Overall Difficulty Analysis")
    lines = [f"Marks for {r} : {m}" for r, m in air]
    h = 0.62 * len(lines) + 0.4
    box = add_box(s, 2.0, max(1.6, (SLIDE_H + HEADER_H - h) / 2), 9.33, h)
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    write_paras(box, lines, size=30, face=TXT_FACE, bold=True, align=PP_ALIGN.CENTER, space_after=4)


# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Build the Paper Analysis deck from a content JSON.")
    ap.add_argument("content", help="content JSON produced by the skill")
    ap.add_argument("-o", "--output", help="output .pptx (default: <content name>.pptx)")
    ap.add_argument("--max-bullets", type=int, default=6, help="max solution bullets per slide (default 6)")
    ap.add_argument("--equations", choices=["native", "latex"], default="native",
                    help="native: real PowerPoint equations, switchable to/from LaTeX inside PowerPoint "
                         "(default; needs pandoc - pip install pypandoc_binary). latex: leave $...$ text.")
    args = ap.parse_args()

    try:
        with open(args.content, encoding="utf-8") as fh:
            data = json.load(fh)
    except (OSError, json.JSONDecodeError) as e:
        sys.exit(f"Could not read {args.content}: {e}")

    qs = data.get("questions") or []
    if not qs:
        sys.exit("The JSON has no questions.")
    need = ["q_no", "type_label", "marks", "concepts", "level", "question", "answer", "solution"]
    for q in qs:
        missing = [k for k in need if k not in q]
        if missing:
            sys.exit(f"Question {q.get('q_no', '?')} is missing: {', '.join(missing)}")

    labels = dict(DEFAULT_LABELS)
    labels.update({str(k): v for k, v in (data.get("level_labels") or {}).items()})

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    warnings = []

    if args.equations == "native":
        pandoc = find_pandoc()
        if not pandoc:
            print("pandoc not found - equations left as LaTeX text. For editable equations run:\n"
                  "    pip install pypandoc_binary")
        else:
            snippets = all_latex(data)
            try:
                EQ["omml"] = keep_safe(convert_latex(snippets, pandoc))
                EQ["mode"] = "native"
            except Exception as e:           # never lose the deck over an equation problem
                print(f"Equation conversion failed ({e}); equations left as LaTeX text.")

    title_slide(prs, data.get("title", "Paper Analysis"), data.get("subtitle", ""))
    starts = {sec["from"]: (i + 1, sec) for i, sec in enumerate(build_sections(data, qs))}
    for q in qs:
        if int(q["q_no"]) in starts:
            section_slide(prs, *starts[int(q["q_no"])])
        info_moved = question_slide(prs, q, labels, warnings)
        question_solution_slides(prs, q, labels, args.max_bullets, warnings, info_moved)
    table_slides(prs, qs, labels)
    summary_slide(prs, qs)
    profile_slide(prs, qs)
    if data.get("air_marks"):
        air_slide(prs, data["air_marks"])
    reference = data.get("reference_indices") or DEFAULT_REFERENCE
    paper_name = data.get("paper_name") or " ".join(
        x for x in [data.get("title", "").split("—")[0].strip(), data.get("subtitle", "")] if x) or "This paper"
    comparison_slide(prs, qs, reference, paper_name)

    wrap_math_shapes()
    out = args.output or os.path.splitext(args.content)[0].replace("_content", "") + "_Analysis.pptx"
    prs.save(out)
    if EQ["mode"] == "native":
        print(f"Equations: {len(EQ['omml'])} converted to native PowerPoint equations.")
    for bad in sorted(EQ["failed"]):
        why = EQ["why"].get(bad)
        warnings.append(f"could not convert ${bad}$ - left as LaTeX text on the slide" + (f" ({why})." if why else "."))
    print(f"Saved {out}  ({len(prs.slides)} slides, {len(qs)} questions)")
    for w in warnings:
        print("  note:", w)


if __name__ == "__main__":
    main()
