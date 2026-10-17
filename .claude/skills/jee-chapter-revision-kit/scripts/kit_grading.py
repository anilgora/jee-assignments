# -*- coding: utf-8 -*-
"""Numbering, question mapping, batch grading, practice-map text, Excel workbook and report."""
import json, re
from collections import defaultdict
from kit_common import pick

SHEET_ORDER = ['J-I', 'J-II', 'J-M', 'J-A', 'C-I', 'C-II']
HIGH_FRACTION = 0.091          # High = tested in about 9.1% of the graded questions (3 of 33 originally)


# ------------------------------------------------------------------ spec / numbering
def load(spec_path, qmap_path, qindex_path):
    spec = json.load(open(spec_path, encoding='utf-8'))
    qmap = json.load(open(qmap_path, encoding='utf-8'))
    qidx = json.load(open(qindex_path, encoding='utf-8'))
    return spec, qmap, qidx


def numbering(spec):
    """assign sequential numbers; returns (items_by_id, newno_by_original_number)"""
    n = 0; byid = {}; newno = {}
    for sec in spec['quick_review']['sections']:
        for it in sec['items']:
            n += 1; it['n'] = n; byid[it['id']] = it
            if it.get('o'): newno[it['o']] = n
    return byid, newno


def xref(t, newno):
    return re.sub(r'\{#(\d+)\}', lambda m: str(newno[int(m.group(1))]), t)


def parse_map(v):
    v = v.strip()
    if v == '-':
        return None, []
    m, _, o = v.partition(';')
    return m.strip(), [s.strip() for s in o.split(',') if s.strip()]


# ------------------------------------------------------------------ grading per batch
def graded_info(spec, qidx, batch):
    jm = spec['batches'][batch]['jm']
    skip = set(spec.get('qb_overlap', {}).get(batch, []))
    sources = spec.get('graded_sources', ['<jm>', 'J-A', 'QB'])
    sources = [jm if s == '<jm>' else s for s in sources]
    n = 0
    for s in sources:
        if s not in qidx: continue
        n += len([q for q in qidx[s] if not (s == 'QB' and q['no'] in skip)])
    return jm, skip, sources, n


def high_threshold(n):
    return max(2, round(HIGH_FRACTION * n))


def grade_of(c, hi):
    return 'High' if c >= hi else ('Medium' if c >= 1 else 'Low')


def batch_context(spec, qmap, qidx, batch):
    jm, skip, gsources, ngr = graded_info(spec, qidx, batch)
    other_jms = {b['jm'] for k, b in spec['batches'].items() if k != batch}
    main = defaultdict(lambda: defaultdict(lambda: defaultdict(list)))
    oth = defaultdict(lambda: defaultdict(lambda: defaultdict(list)))
    count = defaultdict(int)
    for key, v in qmap.items():
        src, no = key.rsplit(':', 1); no = int(no)
        if src in other_jms: continue
        m, o = parse_map(v)
        if src == jm: lab, grp = 'J-M', 'Sheet'
        elif src.startswith('DPP'): lab, grp = src.split('-')[1], 'DPP'
        elif src == 'QB': lab, grp = 'QB', 'QB'
        else: lab, grp = src, 'Sheet'
        if m: main[m][grp][lab].append(no)
        for i in o: oth[i][grp][lab].append(no)
        if src in gsources and not (src == 'QB' and no in skip):
            for i in ([m] if m else []) + o: count[i] += 1
    hi = high_threshold(ngr)
<<<<<<< HEAD
    return dict(batch=batch, main=main, oth=oth, count=count, hi=hi, ngraded=ngr, jm=jm, skip=sorted(skip))
=======
    dpp_only = bool(gsources) and all(s.startswith('DPP') for s in gsources)   # chapter with DPP only (no sheets / QB)
    return dict(batch=batch, main=main, oth=oth, count=count, hi=hi, ngraded=ngr, jm=jm, skip=sorted(skip), dpp_only=dpp_only)
>>>>>>> origin/main


def rng(nums):
    nums = sorted(set(nums)); out = []; i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
        out.append(str(nums[i]) if i == j else f"{nums[i]}-{nums[j]}"); i = j + 1
    return ', '.join(out)


def fmt_groups(groups, lang):
    parts = []
    if groups.get('Sheet'):
        d = groups['Sheet']
        parts.append(('Sheet' if lang == 'en' else 'शीट') + ': ' + '; '.join(f"{ex} {rng(d[ex])}" for ex in SHEET_ORDER if ex in d))
    if groups.get('DPP'):
        d = groups['DPP']; parts.append('DPP: ' + '; '.join(f"{k}({rng(d[k])})" for k in sorted(d, key=int)))
    if groups.get('QB'):
        parts.append(('Question Bank' if lang == 'en' else 'प्रश्न बैंक') + ': ' + rng(groups['QB']['QB']))
    return ' | '.join(parts)


LV = {'en': {'High': 'High', 'Medium': 'Medium', 'Low': 'Low'}, 'hi': {'High': 'उच्च (High)', 'Medium': 'मध्यम (Medium)', 'Low': 'निम्न (Low)'}}
LABELS = {'en': dict(grade='Grade', pm_main='Practice map, main concept: ', pm_oth='Also used in: ', none='none'),
          'hi': dict(grade='श्रेणी', pm_main='अभ्यास मानचित्र, मुख्य अवधारणा: ', pm_oth='इनमें भी प्रयुक्त: ', none='कोई नहीं')}


def item_grade_lines(ctx, item_id, lang):
    g = grade_of(ctx['count'].get(item_id, 0), ctx['hi'])
    L = LABELS[lang]
    ml = fmt_groups(ctx['main'].get(item_id, {}), lang) or L['none']
    ol = fmt_groups(ctx['oth'].get(item_id, {}), lang) or L['none']
    return g, L['pm_main'] + ml, L['pm_oth'] + ol


def legend(ctx, lang):
    ng, hi = ctx['ngraded'], ctx['hi']
<<<<<<< HEAD
=======
    if ctx.get('dpp_only'):
        if lang == 'en':
            return (f"Grade shows how often the item is tested in the DPP questions ({ng} questions): High = {hi} or more questions, Medium = 1 to {hi-1} questions, Low = none. "
                    "Low only means the item is not tested in these questions; it can still be important for understanding. "
                    "Practice map lists the DPP questions that test the item. "
                    "'Main concept' = the question is mainly about this item; 'Also used in' = the question needs it as a secondary step.")
        return (f"श्रेणी दर्शाती है कि DPP के प्रश्नों में ({ng} प्रश्न) कोई बिन्दु कितनी बार पूछा गया है: "
                f"उच्च = {hi} या अधिक प्रश्न, मध्यम = 1 से {hi-1} प्रश्न, निम्न = कोई नहीं। निम्न का अर्थ केवल इतना है कि ये प्रश्न इसे नहीं पूछते; समझ के लिए यह फिर भी महत्त्वपूर्ण हो सकता है। "
                "अभ्यास मानचित्र उन DPP प्रश्नों को सूचीबद्ध करता है जो इस बिन्दु को परखते हैं। "
                "'मुख्य अवधारणा' = प्रश्न मुख्यतः इसी बिन्दु पर है; 'इनमें भी प्रयुक्त' = प्रश्न में यह द्वितीयक चरण है।")
>>>>>>> origin/main
    if lang == 'en':
        return (f"Grade shows how often the item is tested in the JM exercise of this batch, the JA exercise and the JEE Main Question Bank taken together "
                f"({ng} questions; a question that appears in two of these is counted once): High = {hi} or more questions, Medium = 1 to {hi-1} questions, Low = none. "
                "Low only means the item is not tested in these questions; it can still be important for understanding. "
                "Practice map lists the questions that test the item in the sheet exercises (J-I, J-II, J-M, J-A, C-I, C-II), the DPPs and the Question Bank (QB). "
                "'Main concept' = the question is mainly about this item; 'Also used in' = the question needs it as a secondary step.")
    return (f"श्रेणी दर्शाती है कि इस बैच के JM अभ्यास, JA अभ्यास तथा JEE Main प्रश्न बैंक में मिलाकर ({ng} प्रश्न; दो में आने वाला प्रश्न एक बार गिना गया) कोई बिन्दु कितनी बार पूछा गया है: "
            f"उच्च = {hi} या अधिक प्रश्न, मध्यम = 1 से {hi-1} प्रश्न, निम्न = कोई नहीं। निम्न का अर्थ केवल इतना है कि ये प्रश्न इसे नहीं पूछते; समझ के लिए यह फिर भी महत्त्वपूर्ण हो सकता है। "
            "अभ्यास मानचित्र उन प्रश्नों को सूचीबद्ध करता है जो इस बिन्दु को शीट के अभ्यासों (J-I, J-II, J-M, J-A, C-I, C-II), DPP तथा प्रश्न बैंक (QB) में परखते हैं। "
            "'मुख्य अवधारणा' = प्रश्न मुख्यतः इसी बिन्दु पर है; 'इनमें भी प्रयुक्त' = प्रश्न में यह द्वितीयक चरण है।")


# ------------------------------------------------------------------ validation
def validate(spec, qmap, qidx):
    problems = []
    byid, newno = numbering(spec)
    ids = [it['id'] for s in spec['quick_review']['sections'] for it in s['items']]
    if len(ids) != len(set(ids)): problems.append('duplicate item ids')
    def chk(t, where, langs=('en', 'hi')):
        for lang in langs:
            s = pick(t, lang)
            if s.count('$') % 2: problems.append(f'odd number of $ in {where}: {s[:50]}')
            if lang == 'hi' and re.search(r'\\text\{[^}]*[\u0900-\u097F]', s): problems.append(f'Hindi inside \\text in {where}')
            if lang == 'en' and re.search(r'[\u0900-\u097F]', s): problems.append(f'Devanagari in English text in {where}: {s[:40]}')
            for m in re.findall(r'\{#(\d+)\}', s):
                if int(m) not in newno: problems.append(f'unresolved xref {m} in {where}')
    for s in spec['quick_review']['sections']:
        for it in s['items']:
            chk(it['q'], f"item {it['id']} question")
            for l in it['a']: chk(l, f"item {it['id']} answer")
    for s in spec['notes']['sections']:
        for k, ls in s['blocks']:
            for l in ls: chk(l, 'notes ' + s['en'][:20])
    for s in spec['formula']['sections']:
        for l in s['lines']: chk(l, 'formula ' + s['en'][:20])
    for b in spec['mindmap']['branches']:
        for l in b['leaves']: chk(l, 'mindmap ' + b['en'][:20])
    for c in spec['flowcharts']:
        chk(c['start'], 'flowchart start')
        for st in c['steps']:
            chk(st['q'], 'flowchart question')
            for l in st['a'][0]: chk(l, 'flowchart answer (en)', ('en',))
            for l in st['a'][1]: chk(l, 'flowchart answer (hi)', ('hi',))
    for key, v in qmap.items():
        m, o = parse_map(v)
        for i in ([m] if m else []) + o:
            if i not in byid: problems.append(f'question map {key} refers to unknown item id {i}')
    # Hindi lint: English words outside math in Hindi strings
    for s in spec['notes']['sections']:
        for k, ls in s['blocks']:
            for l in ls:
                h = pick(l, 'hi'); w = re.findall(r'[A-Za-z]{3,}', re.sub(r'\$[^$]*\$', '', h))
                if w: problems.append(f'English words in Hindi notes line: {w[:3]}')
    return problems


# ------------------------------------------------------------------ Excel
def build_excel(spec, qmap, qidx, ctxs, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import CellIsRule
    byid, newno = numbering(spec)
    batches = list(ctxs.keys())
    stems = {f"{s}:{q['no']}": q['stem'] for s, qs in qidx.items() for q in qs}
    F = lambda **k: Font(name='Arial', size=10, **k)
    hdr = PatternFill('solid', fgColor='1F3864'); thin = Side(style='thin', color='BFBFBF'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook(); ws = wb.active; ws.title = 'Items'; mp = wb.create_sheet('Map'); un = wb.create_sheet('Unmapped'); st = wb.create_sheet('Settings')
    def head(w, cols):
        w.append(cols)
        for c in w[1]: c.fill = hdr; c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF'); c.alignment = Alignment(wrap_text=True, vertical='center')
    st['A1'] = 'Grade thresholds (edit the blue cells; grades on the Items sheet recalculate)'; st['A1'].font = F(bold=True)
    for i, b in enumerate(batches, 3):
        st[f'A{i}'] = f'Batch {b}: minimum graded questions for High'; st[f'B{i}'] = ctxs[b]['hi']
    medrow = 3 + len(batches)
    st[f'A{medrow}'] = 'Minimum graded questions for Medium (all batches)'; st[f'B{medrow}'] = 1
    for r in range(3, medrow + 1):
        st[f'A{r}'].font = F(); st[f'B{r}'].font = Font(name='Arial', size=10, color='0000FF'); st[f'B{r}'].fill = PatternFill('solid', fgColor='FFFF00')
    notes = ['Graded questions = JM exercise of the batch + JA exercise + JEE Main Question Bank; a question present in two of these counts once.'] + \
            [f"Batch {b}: {ctxs[b]['ngraded']} graded questions; High threshold {ctxs[b]['hi']} (about 9.1%)." for b in batches] + \
            ['Sheet exercises J-I, J-II, Challenger (C-I, C-II) and the DPPs are mapped but not graded; Challenger exercises are not used for the coverage check.',
             'A question counts for every item it needs (main concept and also-used items). Low = not tested in the graded questions; it does not mean unimportant.']
    for i, t in enumerate(notes, medrow + 2): st[f'A{i}'] = t; st[f'A{i}'].font = F(italic=True)
    st.column_dimensions['A'].width = 120; st.column_dimensions['B'].width = 10
    head(mp, ['Source', 'Q no.', 'Group', 'Item no.', 'Role'] + [f'Graded in {b}' for b in batches] + ['Question (English, truncated)'])
    rows = 0
    for key, v in qmap.items():
        src, no = key.rsplit(':', 1); no = int(no); m, o = parse_map(v)
        grp = 'DPP' if src.startswith('DPP') else ('Question Bank' if src == 'QB' else 'Sheet')
        lab = src
        flags = []
        for b in batches:
            c = ctxs[b]
            flags.append('Y' if (src in ('J-A',) or src == c['jm'] or (src == 'QB' and no not in c['skip'])) else 'N')
        for role, ids in (('Main', [m] if m else []), ('Also', o)):
            for i in ids:
                mp.append([lab, no, grp, byid[i]['n'], role] + flags + [stems.get(key, '')[:140]]); rows += 1
    for r in mp.iter_rows(min_row=2):
        for c in r: c.font = F(); c.border = box
    mp.column_dimensions['H'].width = 100; mp.freeze_panes = 'A2'
    head(un, ['Source', 'Q no.', 'Question (English, truncated)'])
    for key, v in qmap.items():
        if parse_map(v)[0] is None:
            s, no = key.rsplit(':', 1); un.append([s, int(no), stems.get(key, '')[:140]])
    for r in un.iter_rows(min_row=2):
        for c in r: c.font = F(); c.border = box
    un.column_dimensions['C'].width = 110
    cols = ['No.', 'Heading', 'Quick Review prompt', 'Status', 'Added from', 'Original no.']
    for b in batches: cols += [f'{b} graded count', f'{b} grade']
    cols += ['All mapped questions']
    head(ws, cols)
    letters = 'ABCDEFGHIJKLMNOP'
    r = 1; gcols = []
    for sec in spec['quick_review']['sections']:
        for it in sec['items']:
            r += 1
            status = 'Added' if it['tag'] == 'added' else ('Rephrased' if it['tag'] == 'rephrased' else 'Original')
            row = [it['n'], sec['en'], re.sub(r'\$', '', it['q'][0]), status, it.get('src', ''), it.get('o')]
            for bi, b in enumerate(batches):
                cc = letters[6 + 2 * bi]; gc = letters[7 + 2 * bi]
                flagcol = letters[5 + bi]            # Map sheet column F.. for this batch
                row += [f'=COUNTIFS(Map!$D:$D,$A{r},Map!${flagcol}:${flagcol},"Y")',
                        f'=IF({cc}{r}>=Settings!$B${3+bi},"High",IF({cc}{r}>=Settings!$B${medrow},"Medium","Low"))']
            row += [f'=COUNTIF(Map!$D:$D,$A{r})']
            ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = F(); c.border = box; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = f'A1:{letters[len(cols)-1]}{r}'
    for col, w in zip('ABCDEF', (6, 30, 70, 12, 24, 10)): ws.column_dimensions[col].width = w
    for bi in range(len(batches)):
        gcol = letters[7 + 2 * bi]
        rngx = f'{gcol}2:{gcol}{r}'
        ws.conditional_formatting.add(rngx, CellIsRule(operator='equal', formula=['"High"'], fill=PatternFill('solid', bgColor='F4B6B6', fgColor='F4B6B6')))
        ws.conditional_formatting.add(rngx, CellIsRule(operator='equal', formula=['"Medium"'], fill=PatternFill('solid', bgColor='FFE0A3', fgColor='FFE0A3')))
        ws.conditional_formatting.add(rngx, CellIsRule(operator='equal', formula=['"Low"'], fill=PatternFill('solid', bgColor='E7E6E6', fgColor='E7E6E6')))
    wb.save(path)


# ------------------------------------------------------------------ report
def build_report(spec, qmap, qidx, ctxs, problems, path):
    byid, newno = numbering(spec)
    L = [f"# {spec['chapter']['en']}: revision kit report", '']
    n_items = len(byid)
    L.append(f'- Items in the Quick Review: {n_items}')
    for b, c in ctxs.items():
        gs = [grade_of(c['count'].get(i, 0), c['hi']) for i in byid]
        L.append(f"- Batch {b}: {c['ngraded']} graded questions, High threshold {c['hi']} -> High {gs.count('High')}, Medium {gs.count('Medium')}, Low {gs.count('Low')}")
        hi_items = [it['n'] for i, it in byid.items() if grade_of(c['count'].get(i, 0), c['hi']) == 'High']
        L.append(f"  - High items (numbers): {sorted(hi_items)}")
    L.append('')
    L.append('## Added items'); 
    for i, it in byid.items():
        if it['tag'] == 'added': L.append(f"- {it['n']}: {it['q'][0]}  [source: {it.get('src','')}]")
    L.append('\n## Rephrased items')
    for i, it in byid.items():
        if it['tag'] == 'rephrased': L.append(f"- {it['n']} (original {it.get('o')}): {it['q'][0]}")
    L.append('\n## Questions with no matching item')
    for key, v in qmap.items():
        if parse_map(v)[0] is None: L.append(f'- {key}')
    L.append('\n## Question Bank questions counted once (overlap with a batch JM set)')
    for b, c in ctxs.items():
        if c['skip']: L.append(f"- Batch {b}: QB {c['skip']}")
    L.append('\n## Validation warnings')
    L += [f'- {p}' for p in problems] or ['- none']
    open(path, 'w', encoding='utf-8').write('\n'.join(L) + '\n')
